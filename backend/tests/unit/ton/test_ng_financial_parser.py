"""DATA-002 parser contracts on synthetic workbooks. No customer data."""

import re
from collections.abc import Sequence
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Any, cast
from uuid import UUID
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from onyx.error_handling.exceptions import OnyxError
from onyx.ton.ng_financial import parser as parser_module
from onyx.ton.ng_financial.diff import compare_source_imports
from onyx.ton.ng_financial.models import (
    DiagnosticCode,
    DiagnosticLevel,
    ParsedImportResult,
    RowKind,
)
from onyx.ton.ng_financial.parser import (
    HEADER,
    NgFinancialExportParser,
    parse_date,
    parse_decimal,
)

SNAPSHOT_ID = UUID(int=1)
DAY_1 = date(2026, 1, 2)
DAY_2 = date(2026, 1, 3)
Row = list[object]


def launch(
    account: str | None = None,
    day: object = None,
    unit: str | None = None,
    document: str | None = None,
    history: str = "001 - Synthetic launch",
    amount: object = 10,
    interest: object = 0,
) -> Row:
    """One A:Q row. Interest is column K; final amount is column Q."""
    return [
        account,
        day,
        unit,
        document,
        history,
        amount,
        0,
        amount,
        None,
        None,
        interest,
        0,
        0,
        0,
        0,
        0,
        amount,
    ]


def book(sheets: dict[str, Sequence[Row]]) -> bytes:
    workbook = Workbook()
    workbook.remove(cast(Worksheet, workbook.active))
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse(content: bytes, snapshot_id: UUID = SNAPSHOT_ID) -> ParsedImportResult:
    return NgFinancialExportParser().parse(content, snapshot_id)


def codes(result: ParsedImportResult, level: DiagnosticLevel) -> list[tuple[str, int]]:
    return [
        (item.code.value, item.row_number or 0)
        for item in result.diagnostics
        if item.level == level
    ]


def tree_rows() -> list[Row]:
    """Root 1 has child 1.1 with leaves 1.1.0001 and 1.1.0002. Root 2 is a
    depth-one leaf. Each launch appears once per ancestor level."""
    a = launch(day=DAY_1, unit="U1", document="D-1", amount=10)
    b = launch(day=DAY_2, unit="U1", document="D-2", amount=20)
    c = launch(day=DAY_2, unit="U2", document="D-3", amount=30)
    rows: list[Row] = []
    for code in ("1 - Root", "1.1 - Group"):
        rows += [
            [code, *a[1:]],
            b,
            [None, None, *c[2:]],
        ]
    rows += [["1.1.0001 - Leaf A", *a[1:]], b]
    rows += [["1.1.0002 - Leaf B", *c[1:]]]
    rows += [["2 - Standalone", *launch(day=DAY_1, unit="U3", amount=5)[1:]]]
    return [list(row) for row in rows]


def test_leaf_selection_prevents_hierarchy_double_counting() -> None:
    result = parse(book({"Jan": tree_rows()}))
    assert [record.locator.row_number for record in result.records] == [7, 8, 9, 10]
    assert {record.account_code for record in result.records} == {
        "1.1.0001",
        "1.1.0002",
        "2",
    }
    summary = result.sheets[0]
    assert summary.classifications[RowKind.HIERARCHY] == 6
    assert summary.classifications[RowKind.DETAIL] == 4
    assert summary.row_spans[RowKind.HIERARCHY] == [(1, 6)]
    # The hierarchy is fully reconciled on F:Q, so nothing is flagged.
    assert codes(result, DiagnosticLevel.WARNING) == []
    assert sum(record.final_amount or 0 for record in result.records) == 65


def test_parent_only_and_child_only_rows_are_reported_not_imported() -> None:
    rows = tree_rows()
    rows.insert(3, launch(amount=99))  # Only in the root block.
    result = parse(book({"Jan": rows}))
    assert len(result.records) == 4
    assert ("PARENT_ROW_WITHOUT_CHILD_MATCH", 4) in codes(
        result, DiagnosticLevel.WARNING
    )
    rows = tree_rows()
    rows.append(launch(amount=7))  # Root-level leaf: nothing to reconcile.
    rows.insert(8, launch(amount=8))  # Leaf A row missing from 1.1.
    result = parse(book({"Jan": rows}))
    assert ("CHILD_ROW_WITHOUT_PARENT_MATCH", 9) in codes(
        result, DiagnosticLevel.WARNING
    )
    assert len(result.records) == 6


def test_grouped_date_and_nested_unit_carry_forward() -> None:
    rows = [
        launch("1 - Root", DAY_1, "U1", "D-1"),
        launch(document="D-2"),
        launch(day=DAY_2, document="D-3"),
        launch(document="D-4"),
        launch(unit="U2", document="D-5"),
        launch("1.1 - Leaf", DAY_1, "U1", "D-1"),
        launch(document="D-2"),
        # A new date without C has no unit; C is not carried across dates.
        launch(day=DAY_2, document="D-3"),
        launch(document="D-4"),
        launch(unit="U2", document="D-5"),
    ]
    result = parse(book({"Fev": rows}))
    got = [
        (record.emission_date, record.administrative_unit, record.source_values["B"])
        for record in result.records
    ]
    assert got == [
        (DAY_1, "U1", "2026-01-02 00:00:00"),
        (DAY_1, "U1", None),
        (DAY_2, None, "2026-01-03 00:00:00"),
        (DAY_2, None, None),
        (DAY_2, "U2", None),
    ]
    assert codes(result, DiagnosticLevel.WARNING) == [
        ("UNIT_BLANK", 8),
        ("UNIT_BLANK", 9),
    ]


def test_invalid_date_is_never_replaced_by_the_previous_date() -> None:
    rows = [
        launch("1 - Root", DAY_1, "U1", "D-1"),
        launch("1.1 - Leaf", DAY_1, "U1", "D-1"),
        launch(None, "31/02/2026", "U1", "D-2"),
        launch(document="D-3"),
        launch(None, DAY_2, "U1", "D-4"),
        launch("1.2 - Leaf", None, "U1", "D-5"),
    ]
    result = parse(book({"Jan": rows}))
    assert [record.locator.row_number for record in result.records] == [2, 5]
    assert codes(result, DiagnosticLevel.ERROR) == [
        ("INVALID_DATE", 3),
        ("INHERITED_DATE_INVALID", 4),
        ("MISSING_DATE", 6),
    ]


def test_malformed_amount_rejects_only_its_row_with_location() -> None:
    rows = tree_rows()
    rows[7][10] = "not a number"
    result = parse(book({"Jan": rows}))
    assert [record.locator.row_number for record in result.records] == [7, 9, 10]
    errors = [item for item in result.diagnostics if item.level == "ERROR"]
    assert [(item.code, item.row_number, item.column) for item in errors] == [
        (DiagnosticCode.INVALID_AMOUNT, 8, "K")
    ]
    # Diagnostics are location-only.
    assert "not a number" not in result.model_dump_json()


def test_blank_and_zero_are_distinct() -> None:
    result = parse(book({"Jan": tree_rows()}))
    record = result.records[0]
    assert record.installment_retention_amount is None
    assert record.source_values["I"] is None
    assert record.interest_amount == Decimal("0")
    assert record.source_values["K"] == "0"
    assert record.document_number == "D-1"


def test_missing_amounts_and_unknown_rows_are_rejected() -> None:
    rows = tree_rows()
    rows.append([None, None, None, "D-9", "History only"])
    rows.append(["free text", None, None, None, None, 1])
    rows.append([None, DAY_1, "U1"])
    result = parse(book({"Jan": rows}))
    assert codes(result, DiagnosticLevel.ERROR) == [
        ("MISSING_AMOUNTS", 11),
        ("UNKNOWN_ROW", 12),
        ("UNKNOWN_ROW", 13),
    ]
    assert len(result.records) == 4


def test_header_section_subtotal_total_and_blank_rows() -> None:
    rows = tree_rows()
    rows[7:7] = [list(HEADER)]
    rows.append(list(HEADER))
    rows.append(["2.1 - Label only"])
    rows.append(["Subtotal", None, None, None, None, 1])
    rows.append([None, None, None, None, "Subtotal da conta", 1])
    rows.append(["Total", None, None, None, None, 1])
    rows.append(["Total geral", None, None, None, None, 1])
    rows.append([])
    rows.append(launch(day=DAY_1, unit="U4", document="D-7"))
    result = parse(book({"Mar": rows}))
    counts = result.sheets[0].classifications
    assert counts[RowKind.HEADER] == 1
    assert counts[RowKind.REPEATED_HEADER] == 1
    assert counts[RowKind.SECTION] == 1
    assert counts[RowKind.SUBTOTAL] == 2
    assert counts[RowKind.TOTAL] == 2
    assert counts[RowKind.BLANK] == 1
    # 2 has a child now, so the standalone launch became hierarchy.
    assert {record.account_code for record in result.records} == {
        "1.1.0001",
        "1.1.0002",
        "2.1",
    }


def test_blank_row_inside_sheet_is_classified() -> None:
    rows = tree_rows()
    spaced = [*rows[:7], [None] * 3, *rows[7:]]
    first, second = parse(book({"Jan": rows})), parse(book({"Jan": spaced}))
    assert second.sheets[0].classifications[RowKind.BLANK] == 1
    # Moving rows changes locators, never fingerprints.
    assert [r.fingerprint for r in first.records] == [
        r.fingerprint for r in second.records
    ]
    assert [r.locator.row_number for r in second.records] == [7, 9, 10, 11]


@pytest.mark.parametrize(
    ("name", "month"),
    [
        ("Jan", 1),
        ("Fev", 2),
        ("Mar", 3),
        ("Abr", 4),
        ("Mai", 5),
        ("Jun", 6),
        ("Jul", 7),
        ("Ago", 8),
        ("Set", 9),
        ("Out", 10),
        ("Nov", 11),
        ("Dez", 12),
        ("Abr ok", 4),
        ("JUN OK", 6),
        ("Mai-ok", 5),
    ],
)
def test_monthly_sheet_discovery(name: str, month: int) -> None:
    result = parse(book({name: tree_rows()}))
    assert result.sheets[0].accepted and result.sheets[0].month == month
    assert {record.sheet_month for record in result.records} == {month}
    assert result.records[0].locator.sheet_name == name


def test_payroll_and_other_sheets_are_skipped_unread() -> None:
    payroll_like: list[Row] = [["Vencimentos", "Ref."], ["=SUM(A1:A2)"]]
    result = parse(
        book(
            {
                "JanFG": payroll_like,
                "Jan-f": payroll_like,
                "Plan1": [],
                "Jan": tree_rows(),
                "Janeiro": tree_rows(),
            }
        )
    )
    skipped = [sheet.name for sheet in result.sheets if not sheet.accepted]
    assert skipped == ["JanFG", "Jan-f", "Plan1", "Janeiro"]
    assert {record.locator.sheet_name for record in result.records} == {"Jan"}
    assert [
        item.sheet_name
        for item in result.diagnostics
        if item.code == DiagnosticCode.SHEET_OUT_OF_SCOPE
    ] == skipped


def test_structural_profile_rejection() -> None:
    with pytest.raises(OnyxError):
        parse(book({"Planilha": [["unrelated", "data"]]}))
    # A month name alone does not match: the rows must have the NG shape.
    with pytest.raises(OnyxError):
        parse(book({"Jan": [["Name", "Value"], ["Alice", 1]]}))
    with pytest.raises(OnyxError):
        parse(book({"Jan": [launch("1 - Flat", DAY_1, "U1")]}))  # No hierarchy.
    with pytest.raises(OnyxError):
        parse(book({"Jan": tree_rows(), "Jan ok": tree_rows()}))
    with pytest.raises(OnyxError):
        parse(b"not a zip archive")


def test_empty_month_sheet_is_skipped_with_warning() -> None:
    result = parse(book({"Jan": [], "Fev": tree_rows()}))
    assert [sheet.accepted for sheet in result.sheets] == [False, True]
    assert result.diagnostics[0].code == DiagnosticCode.MONTH_SHEET_EMPTY


def test_duplicate_looking_rows_are_preserved() -> None:
    same = launch(None, None, None, "D-1")
    rows = [
        launch("1 - Root", DAY_1, "U1", "D-1"),
        same,
        launch("1.1 - Leaf", DAY_1, "U1", "D-1"),
        same,
    ]
    result = parse(book({"Jan": rows}))
    assert len(result.records) == 2
    first, second = result.records
    assert first.fingerprint == second.fingerprint
    assert (first.duplicate_ordinal, second.duplicate_ordinal) == (1, 2)
    assert first.locator != second.locator


def test_lineage_and_deterministic_reparse() -> None:
    content = book({"Jan": tree_rows()})
    first, second = parse(content), parse(content)
    assert first.model_dump() == second.model_dump()
    for record in first.records:
        assert record.locator.snapshot_id == SNAPSHOT_ID
        assert record.locator.sheet_name == "Jan"
        assert record.locator.row_number is not None
        # A generated hash is never presented as a source identifier.
        assert record.locator.source_record_key is None
        assert re.fullmatch(r"[0-9a-f]{64}", record.fingerprint)
    other = parse(content, UUID(int=9))
    assert [r.fingerprint for r in other.records] == [
        r.fingerprint for r in first.records
    ]


def test_formula_rows_are_rejected_not_evaluated() -> None:
    rows = tree_rows()
    rows[6][16] = "=1+1"
    result = parse(book({"Jan": rows}))
    assert [record.locator.row_number for record in result.records] == [8, 9, 10]
    errors = [item for item in result.diagnostics if item.level == "ERROR"]
    assert [(item.code, item.row_number, item.column) for item in errors] == [
        (DiagnosticCode.FORMULA_UNSUPPORTED, 7, "Q")
    ]


def test_cells_outside_contract_are_preserved_and_reported() -> None:
    rows = tree_rows()
    rows[6] = [*rows[6], None, None, 42]
    result = parse(book({"Jan": rows}))
    assert result.records[0].source_values["T"] == "42"
    assert ("CELL_OUTSIDE_CONTRACT", 7) in codes(result, DiagnosticLevel.INFO)
    baseline = parse(book({"Jan": tree_rows()}))
    assert result.records[0].fingerprint != baseline.records[0].fingerprint


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("1.234,56", Decimal("1234.56")),
        ("R$ 1.234,56", Decimal("1234.56")),
        ("-1.234.567,8", Decimal("-1234567.8")),
        ("1234,5", Decimal("1234.5")),
        ("1.500", Decimal("1500")),
        ("0,00", Decimal("0.00")),
        ("  ", None),
        (None, None),
        (0, Decimal("0")),
        (12.5, Decimal("12.5")),
        (Decimal("3.10"), Decimal("3.10")),
    ],
)
def test_brazilian_amounts(value: object, expected: Decimal | None) -> None:
    assert parse_decimal(value) == expected


@pytest.mark.parametrize(
    "value",
    [
        "1,234.56",
        "1.5",
        "1.2345,6",
        "1 234,56",
        "(10,00)",
        "NaN",
        "bad",
        True,
        float("inf"),
        datetime(2026, 1, 1),
        time(1, 0),
        Decimal("1E+25"),
        Decimal("0.00000000001"),
    ],
)
def test_ambiguous_or_invalid_amounts(value: object) -> None:
    with pytest.raises(ValueError):
        parse_decimal(value)


def test_dates_without_guessing() -> None:
    assert parse_date("31/12/2026") == date(2026, 12, 31)
    assert parse_date("2026-12-31") == date(2026, 12, 31)
    assert parse_date(datetime(2026, 12, 31)) == date(2026, 12, 31)
    assert parse_date(None) is None
    for value in (
        "12/31/2026",
        "31/02/2026",
        "31-12-2026",
        46000,
        datetime(2026, 1, 1, 10, 30),
        time(0, 0),
    ):
        with pytest.raises(ValueError):
            parse_date(value)


def test_workbook_loading_is_offline_and_formula_free(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[dict[str, Any]] = []
    original = parser_module.load_workbook

    def spy(*args: Any, **kwargs: Any) -> Any:
        calls.append(kwargs)
        return original(*args, **kwargs)

    monkeypatch.setattr(parser_module, "load_workbook", spy)
    parse(book({"Jan": tree_rows()}))
    assert calls == [
        {
            "read_only": True,
            "data_only": False,
            "keep_links": False,
            "keep_vba": False,
        }
    ]


def test_macro_container_is_rejected() -> None:
    payload = BytesIO(book({"Jan": tree_rows()}))
    with ZipFile(payload, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"inert synthetic marker")
    with pytest.raises(OnyxError):
        parse(payload.getvalue())


def test_declared_dimension_is_not_trusted() -> None:
    source = BytesIO(book({"Jan": tree_rows()}))
    target = BytesIO()
    with ZipFile(source) as original, ZipFile(target, "w") as rewritten:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(
                    rb'<dimension ref="[^"]+"', b'<dimension ref="A1:B2"', data
                )
            rewritten.writestr(item, data)
    result = parse(target.getvalue())
    assert len(result.records) == 4
    assert result.records[-1].source_values["Q"] == "5"


@pytest.mark.parametrize(
    ("name", "value"),
    [
        ("MAX_ROWS_PER_SHEET", 5),
        ("MAX_TOTAL_ROWS", 5),
        ("MAX_SHEETS", 1),
        ("MAX_WORKBOOK_BYTES", 10),
        ("MAX_COLUMNS", 10),
        ("MAX_CELL_CHARS", 5),
        ("MAX_SECONDS", -1),
    ],
)
def test_resource_limits(
    monkeypatch: pytest.MonkeyPatch, name: str, value: int
) -> None:
    content = book({"Jan": tree_rows(), "JanFG": [["x"]]})
    monkeypatch.setattr(parser_module, name, value)
    with pytest.raises(OnyxError):
        parse(content)


def month_result(rows: list[Row], snapshot: int) -> ParsedImportResult:
    return parse(book({"Jun": rows}), UUID(int=snapshot))


def test_diff_identical_exports_are_unchanged() -> None:
    diff = compare_source_imports(
        month_result(tree_rows(), 1), month_result(tree_rows(), 2)
    )
    assert (diff.unchanged, diff.changed, len(diff.added), len(diff.removed)) == (
        4,
        0,
        0,
        0,
    )


def test_diff_treats_mutable_field_changes_as_changed_records() -> None:
    reviewed = tree_rows()
    reviewed[6][10] = 3  # Interest.
    reviewed[6][16] = 13  # Final amount.
    reviewed[8][2] = "U9"  # Unit (classification) and history.
    reviewed[8][4] = "001 - Reviewed text"
    diff = compare_source_imports(
        month_result(tree_rows(), 1), month_result(reviewed, 2)
    )
    assert (diff.unchanged, diff.changed, len(diff.added), len(diff.removed)) == (
        2,
        2,
        0,
        0,
    )
    assert [pair.fields for pair in diff.changed_pairs] == [
        ("interest_amount", "final_amount"),
        ("administrative_unit", "history"),
    ]
    assert diff.changed_fields["final_amount"] == 1


def test_diff_inserted_row_is_added_and_shift_is_ignored() -> None:
    reviewed = tree_rows()
    reviewed.insert(8, launch(None, DAY_2, "U7", "D-8", "002 - New", 40))
    diff = compare_source_imports(
        month_result(tree_rows(), 1), month_result(reviewed, 2)
    )
    assert (diff.unchanged, diff.changed, len(diff.removed)) == (4, 0, 0)
    assert [(item.month, item.row_number) for item in diff.added] == [(6, 9)]
    removed = compare_source_imports(
        month_result(reviewed, 2), month_result(tree_rows(), 1)
    )
    assert [(item.month, item.row_number) for item in removed.removed] == [(6, 9)]


def test_diff_pairs_exact_duplicates_and_flags_ambiguity() -> None:
    base = [
        launch("1 - Root", DAY_1, "U1", "D-1"),
        launch(document="D-1"),
        launch("1.1 - Leaf", DAY_1, "U1", "D-1"),
        launch(document="D-1"),
    ]
    reviewed = [list(row) for row in base]
    reviewed[3][16] = 11
    diff = compare_source_imports(month_result(base, 1), month_result(reviewed, 2))
    assert (diff.unchanged, diff.changed) == (1, 1)
    assert diff.changed_pairs[0].right.row_number == 4
    both = [list(row) for row in base]
    both[2][16] = 11
    both[3][16] = 12
    ambiguous = compare_source_imports(month_result(base, 1), month_result(both, 2))
    assert ambiguous.changed == 2 and ambiguous.ambiguous == 2
