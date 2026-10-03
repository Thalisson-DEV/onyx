"""Synthetic DRE persistence, readiness, ACL, and immutable revisions."""

import datetime
import io
import time
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import select, text
from sqlalchemy.exc import DBAPIError
from sqlalchemy.orm import Session

from onyx.db.models import User
from onyx.db.ton import dre as repository
from onyx.db.ton.models import DreCalculationRun, DreResultLine, TonAuditEvent
from onyx.error_handling.exceptions import OnyxError
from onyx.file_store.file_store import FileStore
from onyx.server.ton.dre import calculate as calculate_route
from onyx.server.ton.dre import export_calculation, export_calculation_xlsx
from onyx.ton.dre.models import (
    DreAccountAssignment,
    DreLineDefinition,
    DreLineType,
    DreOperation,
    DreScope,
    DreStructureCreate,
    DreVersionCreate,
)
from tests.external_dependency_unit.ton import factories
from tests.external_dependency_unit.ton.test_financial_domain import (
    build_complete_synthetic_scope,
)

pytest_plugins = ("tests.external_dependency_unit.ton.test_financial_domain",)


def _lines() -> list[DreLineDefinition]:
    return [
        DreLineDefinition(
            code="gross",
            label="Gross revenue",
            position=1,
            line_type=DreLineType.SUBTOTAL,
            operation=DreOperation.SUM_CHILDREN,
        ),
        DreLineDefinition(
            code="service",
            label="Service",
            position=2,
            parent_code="gross",
            line_type=DreLineType.SOURCE_SUM,
        ),
        DreLineDefinition(
            code="cost",
            label="Cost",
            position=3,
            line_type=DreLineType.SOURCE_SUM,
        ),
        DreLineDefinition(
            code="result",
            label="Result",
            position=4,
            line_type=DreLineType.RESULT,
            operation=DreOperation.SUBTRACT,
            operands=["gross", "cost"],
        ),
    ]


def test_synthetic_ready_revision_and_mapping_governance(
    ton_session: Session, admin: User, store: FileStore
) -> None:
    normalization_id, account_id, unit_id = build_complete_synthetic_scope(
        ton_session, admin, store
    )
    first_version = repository.create_structure(
        ton_session,
        admin,
        DreStructureCreate(
            key="synthetic-dre",
            label="Synthetic DRE",
            lines=_lines(),
        ),
    )
    ton_session.commit()
    assert (
        repository.get_structure(
            ton_session, admin, first_version.structure_id
        ).latest_version
        == 1
    )
    scope = DreScope(
        normalization_run_id=normalization_id,
        structure_version_id=first_version.id,
        period=datetime.date(2026, 1, 1),
        unit_id=unit_id,
    )
    initial = repository.readiness(ton_session, admin, scope)
    assert initial.status == "NOT_READY"
    assert initial.blockers["DRE_ACCOUNT_UNMAPPED"] >= 1
    blocked = repository.execute(ton_session, admin, scope)
    assert blocked.status == "NOT_READY"
    assert repository.list_result_lines(ton_session, admin, blocked.id, 10, 0) == []

    pending = repository.create_version(
        ton_session,
        admin,
        first_version.structure_id,
        DreVersionCreate(
            lines=_lines(),
            assignments=[
                DreAccountAssignment(
                    account_id=account_id,
                    line_code="service",
                    status="PENDING_APPROVAL",
                )
            ],
        ),
    )
    ton_session.commit()
    scope.structure_version_id = pending.id
    assert (
        repository.readiness(ton_session, admin, scope).blockers[
            "DRE_MAPPING_PENDING_APPROVAL"
        ]
        >= 1
    )

    approved = repository.create_version(
        ton_session,
        admin,
        first_version.structure_id,
        DreVersionCreate(
            lines=_lines(),
            assignments=[
                DreAccountAssignment(
                    account_id=account_id, line_code="service", status="APPROVED"
                )
            ],
        ),
    )
    ton_session.commit()
    scope.structure_version_id = approved.id
    assert repository.readiness(ton_session, admin, scope).status == "READY"
    started = time.perf_counter()
    ready = repository.execute(ton_session, admin, scope)
    print("synthetic_dre_execute_ms", round((time.perf_counter() - started) * 1000))
    assert ready.status == "READY"
    assert repository.execute(ton_session, admin, scope).id == ready.id
    lines = repository.list_result_lines(ton_session, admin, ready.id, 10, 0)
    by_code = {line.code: line for line in lines}
    assert len(lines) == 4
    assert by_code["service"].realizado == Decimal("100.25")
    assert by_code["gross"].realizado == by_code["service"].realizado
    assert by_code["result"].realizado == Decimal("100.25")
    assert by_code["result"].orcado == by_code["service"].orcado
    assert by_code["result"].variance == (
        by_code["result"].realizado - by_code["result"].orcado
    )
    assert by_code["service"].realizado_ytd == Decimal("100.25")
    assert ready.provenance["normalization_run_id"] == str(normalization_id)
    assert ready.provenance["dre_structure_version_id"] == str(approved.id)
    assert ready.provenance["budget_execution_ids"]
    assert ready.provenance["ng_source_snapshot_id"]
    assert ready.id != blocked.id

    statement = repository.get_statement(ton_session, admin, ready.id)
    assert statement.version.id == approved.id
    assert statement.version.lines[0].line_type == DreLineType.SUBTOTAL
    assert statement.lines == lines
    revisions = repository.list_calculations(
        ton_session, admin, datetime.date(2026, 1, 1), unit_id, 10, 0
    )
    assert {item.id for item in revisions} == {blocked.id, ready.id}
    series = repository.get_period_series(
        ton_session, admin, normalization_id, approved.id, 2026, unit_id, "result"
    )
    assert len(series) == 1
    assert series[0].realizado == by_code["result"].realizado
    actual_page = repository.list_contributors(
        ton_session, admin, ready.id, "service", "ACTUAL", 1, 0
    )
    budget_page = repository.list_contributors(
        ton_session, admin, ready.id, "service", "BUDGET", 1, 0
    )
    assert actual_page.total == 1
    assert actual_page.rows[0].amount == by_code["service"].realizado
    assert actual_page.rows[0].amount_basis == "MOVEMENT"
    assert actual_page.rows[0].sheet_name
    assert budget_page.total == 2
    budget_next_page = repository.list_contributors(
        ton_session, admin, ready.id, "service", "BUDGET", 1, 1
    )
    assert (
        budget_page.rows[0].amount + budget_next_page.rows[0].amount
        == by_code["service"].orcado
    )
    assert (
        repository.list_contributors(
            ton_session, admin, ready.id, "service", "ACTUAL", 1, 1
        ).rows
        == []
    )
    csv_response = export_calculation(ready.id, admin, ton_session)
    assert b"result" in csv_response.body
    assert str(by_code["result"].realizado).encode() in csv_response.body
    assert str(ready.id).encode() in csv_response.body
    with pytest.raises(OnyxError):
        export_calculation(blocked.id, admin, ton_session)

    renamed_lines = _lines()
    renamed_lines[1].label = "Renamed service"
    renamed = repository.create_version(
        ton_session,
        admin,
        first_version.structure_id,
        DreVersionCreate(
            lines=renamed_lines,
            assignments=[
                DreAccountAssignment(
                    account_id=account_id, line_code="service", status="APPROVED"
                )
            ],
        ),
    )
    ton_session.commit()
    scope.structure_version_id = renamed.id
    renamed_run = repository.execute(ton_session, admin, scope)
    assert renamed_run.id != ready.id
    assert (
        repository.get_structure(
            ton_session, admin, first_version.structure_id
        ).latest_version
        == renamed.number
    )
    assert (
        repository.list_result_lines(ton_session, admin, ready.id, 10, 0)[1].label
        == "Service"
    )
    assert (
        repository.list_result_lines(ton_session, admin, renamed_run.id, 10, 0)[1].label
        == "Renamed service"
    )

    zero_ratio_lines = _lines()
    zero_ratio_lines.append(
        DreLineDefinition(
            code="ratio",
            label="Ratio",
            position=5,
            line_type=DreLineType.PERCENTAGE,
            operation=DreOperation.RATIO,
            operands=["gross", "cost"],
        )
    )
    zero_ratio = repository.create_version(
        ton_session,
        admin,
        first_version.structure_id,
        DreVersionCreate(
            lines=zero_ratio_lines,
            assignments=[
                DreAccountAssignment(
                    account_id=account_id, line_code="service", status="APPROVED"
                )
            ],
        ),
    )
    ton_session.commit()
    scope.structure_version_id = zero_ratio.id
    assert (
        repository.readiness(ton_session, admin, scope).blockers[
            "FORMULA_DENOMINATOR_ZERO"
        ]
        == 1
    )
    assert repository.execute(ton_session, admin, scope).status == "NOT_READY"

    with pytest.raises(OnyxError):
        calculate_route(
            DreScope(
                normalization_run_id=normalization_id,
                structure_version_id=approved.id,
                period=datetime.date(2026, 1, 2),
                unit_id=unit_id,
            ),
            admin,
            ton_session,
        )

    audit = list(
        ton_session.scalars(
            select(TonAuditEvent).where(TonAuditEvent.action.like("ton_dre.%"))
        )
    )
    assert audit
    assert any(item.action == "ton_dre.calculate_fail" for item in audit)
    assert all(item.before_state is None and item.after_state is None for item in audit)
    outsider = factories.make_user(ton_session)
    ton_session.commit()
    with pytest.raises(OnyxError):
        repository.get_calculation(ton_session, outsider, ready.id)
    with pytest.raises(DBAPIError):
        ton_session.execute(
            text("UPDATE ton_dre_result_line SET realizado = 0 WHERE run_id = :id"),
            {"id": ready.id},
        )
    ton_session.rollback()
    assert (
        ton_session.scalar(
            select(DreResultLine).where(DreResultLine.run_id == ready.id)
        )
        is not None
    )
    assert ton_session.get(DreCalculationRun, blocked.id) is not None


def test_synthetic_two_period_ytd_and_series(
    ton_session: Session, admin: User, store: FileStore
) -> None:
    normalization_id, account_id, unit_id = build_complete_synthetic_scope(
        ton_session, admin, store, include_february=True
    )
    version = repository.create_structure(
        ton_session,
        admin,
        DreStructureCreate(
            key="synthetic-two-period-dre",
            label="Synthetic two period DRE",
            lines=_lines(),
            assignments=[
                DreAccountAssignment(
                    account_id=account_id, line_code="service", status="APPROVED"
                )
            ],
        ),
    )
    ton_session.commit()
    january = repository.execute(
        ton_session,
        admin,
        DreScope(
            normalization_run_id=normalization_id,
            structure_version_id=version.id,
            period=datetime.date(2026, 1, 1),
            unit_id=unit_id,
        ),
    )
    february = repository.execute(
        ton_session,
        admin,
        DreScope(
            normalization_run_id=normalization_id,
            structure_version_id=version.id,
            period=datetime.date(2026, 2, 1),
            unit_id=unit_id,
        ),
    )
    assert january.status == february.status == "READY"
    january_lines = {
        line.code: line
        for line in repository.list_result_lines(ton_session, admin, january.id, 10, 0)
    }
    february_lines = {
        line.code: line
        for line in repository.list_result_lines(ton_session, admin, february.id, 10, 0)
    }
    assert january_lines["result"].realizado == Decimal("100.25")
    assert february_lines["result"].realizado == Decimal("40.75")
    assert february_lines["result"].realizado_ytd == Decimal("141.00")
    assert february_lines["result"].orcado_ytd == Decimal("20")
    assert february_lines["result"].variance_ytd == Decimal("121.00")
    series = repository.get_period_series(
        ton_session, admin, normalization_id, version.id, 2026, unit_id, "result"
    )
    assert [point.period for point in series] == [
        datetime.date(2026, 1, 1),
        datetime.date(2026, 2, 1),
    ]
    assert [point.realizado for point in series] == [
        Decimal("100.25"),
        Decimal("40.75"),
    ]
    february_actuals = repository.list_contributors(
        ton_session, admin, february.id, "service", "ACTUAL", 10, 0
    )
    assert february_actuals.total == 1
    assert february_actuals.rows[0].period == datetime.date(2026, 2, 1)


def test_invalid_structure_dependencies_rejected(
    ton_session: Session, admin: User
) -> None:
    with pytest.raises(OnyxError):
        repository.create_structure(
            ton_session,
            admin,
            DreStructureCreate(
                key="invalid",
                label="Invalid",
                lines=[
                    DreLineDefinition(
                        code="total",
                        label="Total",
                        position=1,
                        line_type=DreLineType.RESULT,
                        operation=DreOperation.SUM_LINES,
                        operands=["missing"],
                    )
                ],
            ),
        )


def test_synthetic_excel_export_matches_persisted_dre(
    ton_session: Session, admin: User, store: FileStore
) -> None:
    normalization_id, account_id, unit_id = build_complete_synthetic_scope(
        ton_session, admin, store, include_february=True
    )
    version = repository.create_structure(
        ton_session,
        admin,
        DreStructureCreate(
            key="synthetic-excel-dre",
            label="Synthetic Excel DRE",
            lines=_lines(),
            assignments=[
                DreAccountAssignment(
                    account_id=account_id, line_code="service", status="APPROVED"
                )
            ],
        ),
    )
    ton_session.commit()
    runs = {
        (unit, month): repository.execute(
            ton_session,
            admin,
            DreScope(
                normalization_run_id=normalization_id,
                structure_version_id=version.id,
                period=datetime.date(2026, month, 1),
                unit_id=unit,
            ),
        )
        for unit in (None, unit_id)
        for month in (1, 2)
    }
    assert all(run.status == "READY" for run in runs.values())

    response = export_calculation_xlsx(runs[(None, 2)].id, admin, ton_session)
    assert response.headers["content-disposition"].endswith(
        'filename="dre-2026-01-a-02.xlsx"'
    )
    workbook = openpyxl.load_workbook(io.BytesIO(response.body), data_only=True)
    formulas = openpyxl.load_workbook(io.BytesIO(response.body))
    assert workbook.sheetnames == ["DRE", "Base", "Premissas"]
    base = list(workbook["Base"].iter_rows(min_row=2, values_only=True))
    assert sorted(float(str(row[3])) for row in base) == [40.75, 100.25]
    assert {row[2] for row in base} == {"service"}
    assert all(row[15] for row in base)

    sheet = workbook["DRE"]
    blocks = [
        row
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row + 2, 1).value == "Código"
    ]
    # Consolidated block first, then the one synthetic unit.
    assert len(blocks) == 2
    for block, unit in zip(blocks, (None, unit_id), strict=True):
        rows = {
            str(sheet.cell(row, 1).value): row
            for row in range(block + 3, block + 3 + len(_lines()))
        }
        for month, column in ((1, 3), (2, 4)):
            persisted = {
                line.code: line
                for line in repository.list_result_lines(
                    ton_session, admin, runs[(unit, month)].id, 10, 0
                )
            }
            for code, row in rows.items():
                assert Decimal(str(sheet.cell(row, column).value)) == (
                    persisted[code].realizado
                )
                if month == 2:
                    assert Decimal(str(sheet.cell(row, 5).value)) == (
                        persisted[code].realizado_ytd
                    )
        assert str(formulas["DRE"].cell(rows["service"], 3).value).startswith(
            "=SUMIFS(Base[Valor],"
        )
        assert formulas["DRE"].cell(rows["result"], 4).value == (
            f"=D{rows['gross']}-D{rows['cost']}"
        )
    topics = [
        row[0] for row in workbook["Premissas"].iter_rows(values_only=True) if row
    ]
    assert "Parcelamentos" in topics and "PIS/COFINS" in topics
    assert any(
        item.action == "ton_dre.export" and item.resource_id == runs[(None, 2)].id
        for item in ton_session.scalars(select(TonAuditEvent))
    )

    outsider = factories.make_user(ton_session)
    ton_session.commit()
    with pytest.raises(OnyxError):
        export_calculation_xlsx(runs[(None, 2)].id, outsider, ton_session)
    blocked = repository.execute(
        ton_session,
        admin,
        DreScope(
            normalization_run_id=normalization_id,
            structure_version_id=repository.create_version(
                ton_session,
                admin,
                version.structure_id,
                DreVersionCreate(lines=_lines()),
            ).id,
            period=datetime.date(2026, 2, 1),
        ),
    )
    ton_session.commit()
    assert blocked.status == "NOT_READY"
    with pytest.raises(OnyxError):
        export_calculation_xlsx(blocked.id, admin, ton_session)
