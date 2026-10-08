"""Data steps: bring data into the flow (text, JSON, tables, files) and
reshape what earlier steps produced. Pure functions over JSON values; only
"Inserir dados" with a file reads the database."""

import csv
import io
import json
from collections import defaultdict
from decimal import Decimal
from html import escape
from typing import Any

from onyx.ton.automations.conditions import evaluate_group, parse_group
from onyx.ton.automations.expressions import (
    ExpressionError,
    format_money,
    to_number,
    to_text,
)
from onyx.ton.automations.registry import (
    ActionContext,
    ActionError,
    FieldSpec,
    NodeSpec,
    OutputSpec,
    ParamSpec,
    register,
)

MAX_ROWS = 5000
MAX_TEXT = 200_000


def _items(value: Any, label: str = "Lista") -> list[Any]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return value
    if isinstance(value, dict):
        return [value]
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.startswith("["):
            try:
                parsed = json.loads(stripped)
            except json.JSONDecodeError:
                raise ActionError(f"{label}: o texto não é uma lista JSON") from None
            if isinstance(parsed, list):
                return parsed
    raise ActionError(f"{label}: esperava uma lista (ex.: {{{{ steps.x.outputs.items }}}})")


def _field(item: Any, key: str) -> Any:
    if isinstance(item, dict):
        if key in item:
            return item[key]
        current: Any = item
        for part in key.split("."):
            current = current.get(part) if isinstance(current, dict) else None
        return current
    return item if not key else None


# -- Inserir dados -------------------------------------------------------------


def _table_from_text(text: str, has_header: bool) -> tuple[list[str], list[dict[str, Any]]]:
    sample = text[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    return _rows_to_dicts(rows, has_header)


def _rows_to_dicts(rows: list[list[Any]], has_header: bool) -> tuple[list[str], list[dict[str, Any]]]:
    if not rows:
        return [], []
    width = max(len(row) for row in rows)
    if has_header:
        header = [to_text(cell).strip() or f"coluna_{i + 1}" for i, cell in enumerate(rows[0])]
        header += [f"coluna_{i + 1}" for i in range(len(header), width)]
        body = rows[1:]
    else:
        header = [f"coluna_{i + 1}" for i in range(width)]
        body = rows
    if len(body) > MAX_ROWS:
        raise ActionError(f"No máximo {MAX_ROWS} linhas")
    return header, [
        {header[i]: (row[i] if i < len(row) else None) for i in range(width)} for row in body
    ]


def _file_payload(ctx: ActionContext, file_ref: Any, has_header: bool) -> dict[str, Any]:
    from onyx.db.ton import automations as repository

    session, _owner = ctx.require_db()
    file_id = file_ref.get("id") if isinstance(file_ref, dict) else file_ref
    if not file_id:
        raise ActionError("Escolha um arquivo")
    stored = repository.get_file(session, str(file_id))
    if stored is None:
        raise ActionError("Arquivo não encontrado")
    name = stored.name.lower()
    data = bytes(stored.data)
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True) if any(cell is not None for cell in row)]
        columns, records = _rows_to_dicts(rows[: MAX_ROWS + 1], has_header)
        text = "\n".join(";".join(to_text(cell) for cell in row) for row in rows[:200])
        return {"text": text[:MAX_TEXT], "rows": records, "columns": columns, "count": len(records), "data": records, "file_name": stored.name}
    if name.endswith((".csv", ".txt", ".json", ".md")):
        text = data.decode("utf-8-sig", errors="replace")
        if name.endswith(".json"):
            try:
                parsed = json.loads(text)
            except json.JSONDecodeError:
                raise ActionError("O arquivo JSON é inválido") from None
            rows = parsed if isinstance(parsed, list) else []
            return {"text": text[:MAX_TEXT], "data": parsed, "rows": rows, "columns": sorted({k for r in rows if isinstance(r, dict) for k in r}), "count": len(rows), "file_name": stored.name}
        if name.endswith(".csv"):
            columns, records = _table_from_text(text, has_header)
            return {"text": text[:MAX_TEXT], "rows": records, "columns": columns, "count": len(records), "data": records, "file_name": stored.name}
        return {"text": text[:MAX_TEXT], "rows": [], "columns": [], "count": 0, "data": text[:MAX_TEXT], "file_name": stored.name}
    from onyx.file_processing.extract_file_text import extract_file_text

    try:
        text = extract_file_text(io.BytesIO(data), stored.name, break_on_unprocessable=False)
    except Exception as error:  # noqa: BLE001 - parser errors become a step failure
        raise ActionError(f"Não consegui ler o arquivo: {type(error).__name__}") from None
    return {"text": text[:MAX_TEXT], "rows": [], "columns": [], "count": 0, "data": text[:MAX_TEXT], "file_name": stored.name}


def _input(ctx: ActionContext, params: dict[str, Any]) -> dict[str, Any]:
    source = params.get("source") or "text"
    has_header = params.get("has_header", True) is not False
    if source == "file":
        return _file_payload(ctx, params.get("file"), has_header)
    content = params.get("content")
    if source == "json":
        if isinstance(content, (dict, list)):
            parsed = content
        else:
            try:
                parsed = json.loads(to_text(content) or "null")
            except json.JSONDecodeError as error:
                raise ActionError(f"JSON inválido: linha {error.lineno}, coluna {error.colno}") from None
        rows = parsed if isinstance(parsed, list) else []
        columns = sorted({key for row in rows if isinstance(row, dict) for key in row})
        return {"data": parsed, "rows": rows, "columns": columns, "count": len(rows), "text": json.dumps(parsed, ensure_ascii=False)[:MAX_TEXT]}
    text = to_text(content)
    if len(text) > MAX_TEXT:
        raise ActionError("Texto longo demais (máximo 200 mil caracteres)")
    if source == "table":
        columns, rows = _table_from_text(text, has_header)
        return {"data": rows, "rows": rows, "columns": columns, "count": len(rows), "text": text}
    return {"data": text, "rows": [], "columns": [], "count": 0, "text": text}


register(
    NodeSpec(
        type="data.input",
        group="data",
        label="Inserir dados",
        description="Traz dados para o fluxo: texto livre, JSON, tabela colada (CSV) ou arquivo (Excel, CSV, PDF, Word, texto).",
        icon="database",
        params=(
            ParamSpec("source", "Tipo de dado", "select", default="text", dynamic=False, options=(("text", "Texto livre"), ("table", "Tabela (CSV colado)"), ("json", "JSON"), ("file", "Arquivo"))),
            ParamSpec("content", "Conteúdo", "textarea", placeholder="Cole aqui, ou use conteúdo dinâmico", show_if=("source", ("text", "table", "json"))),
            ParamSpec("file", "Arquivo", "file", dynamic=False, show_if=("source", ("file",)), help="Excel, CSV, JSON, TXT, PDF ou Word, até 5 MB."),
            ParamSpec("has_header", "A primeira linha é o cabeçalho", "boolean", default=True, dynamic=False, show_if=("source", ("table", "file"))),
        ),
        outputs=(
            OutputSpec("text", "Texto", "string"),
            OutputSpec("rows", "Linhas (tabela)", "array"),
            OutputSpec("columns", "Colunas", "array"),
            OutputSpec("count", "Quantidade de linhas", "number"),
            OutputSpec("data", "Dados (como vieram)", "any"),
        ),
        executor=_input,
        satisfies=(),
        keywords=("upload", "arquivo", "planilha", "excel", "csv", "colar", "texto", "pdf"),
    )
)


# -- Compor / JSON ---------------------------------------------------------------


register(
    NodeSpec(
        type="data.compose",
        group="data",
        label="Compor",
        description="Monta um valor (texto, número, lista ou objeto) a partir de outros dados.",
        icon="compose",
        params=(ParamSpec("value", "Valor", "textarea", required=True, placeholder="Escreva o texto e use ⚡ para incluir dados"),),
        outputs=(OutputSpec("value", "Valor", "any"),),
        executor=lambda ctx, params: {"value": params.get("value")},
        keywords=("compose", "montar", "texto", "fórmula", "expressão"),
    )
)


def _parse_json(ctx: ActionContext, params: dict[str, Any]) -> dict[str, Any]:
    content = params.get("content")
    if isinstance(content, (dict, list)):
        return {"value": content}
    text = to_text(content).strip()
    start = min([i for i in (text.find("{"), text.find("[")) if i >= 0], default=-1)
    if start < 0:
        raise ActionError("Não há JSON no conteúdo")
    try:
        return {"value": json.loads(text[start:])}
    except json.JSONDecodeError:
        end = max(text.rfind("}"), text.rfind("]"))
        try:
            return {"value": json.loads(text[start : end + 1])}
        except json.JSONDecodeError as error:
            raise ActionError(f"JSON inválido: {error.msg}") from None


register(
    NodeSpec(
        type="data.parse_json",
        group="data",
        label="Ler JSON",
        description="Transforma um texto JSON em dados que os próximos passos podem usar.",
        icon="braces",
        params=(ParamSpec("content", "Conteúdo", "textarea", required=True),),
        outputs=(OutputSpec("value", "Dados", "any"),),
        executor=_parse_json,
        keywords=("parse", "json", "converter"),
    )
)


# -- Lists -----------------------------------------------------------------------


def _filter(ctx: ActionContext, params: dict[str, Any]) -> dict[str, Any]:
    items = _items(ctx.resolve(params.get("items"), None))
    group = parse_group(params.get("condition"))
    kept: list[Any] = []
    try:
        for item in items:
            if evaluate_group(group, {**ctx.scope, "item": item}):
                kept.append(item)
    except ExpressionError as error:
        raise ActionError(f"Regras: {error}") from None
    rest = [item for item in items if item not in kept] if params.get("keep_rest") else []
    total = sum((Decimal(str(to_number(_field(i, "valor")) or 0)) for i in kept), Decimal(0))
    return {"items": kept, "count": len(kept), "rest": rest, "rest_count": len(rest), "total": float(total), "total_formatado": format_money(total)}


register(
    NodeSpec(
        type="data.filter",
        group="data",
        label="Filtrar lista",
        description="Mantém só os itens que atendem às regras. Nas regras, use {{ item.campo }}.",
        icon="filter",
        params=(
            ParamSpec("items", "Lista", "expression", required=True, resolve=False, placeholder="Use ⚡ para escolher a lista"),
            ParamSpec("condition", "Manter quando", "condition", required=True, default={"op": "and", "rules": []}, resolve=False),
            ParamSpec("keep_rest", "Guardar também os que não atendem", "boolean", default=False, dynamic=False),
        ),
        outputs=(
            OutputSpec("items", "Itens que atendem", "array"),
            OutputSpec("count", "Quantidade", "number"),
            OutputSpec("total", "Soma do campo valor", "number"),
            OutputSpec("total_formatado", "Soma do campo valor (R$)", "string"),
            OutputSpec("rest", "Itens que não atendem", "array"),
            OutputSpec("rest_count", "Quantidade que não atende", "number"),
        ),
        executor=_filter,
        keywords=("filter", "where", "onde", "separar"),
    )
)


def _select(ctx: ActionContext, params: dict[str, Any]) -> dict[str, Any]:
    items = _items(ctx.resolve(params.get("items"), None))
    mapping = [entry for entry in params.get("mapping") or [] if str(entry.get("key") or "").strip()]
    if not mapping:
        raise ActionError("Informe ao menos um campo")
    result = []
    for item in items:
        result.append({str(entry["key"]).strip(): ctx.resolve(entry.get("value"), {"item": item}) for entry in mapping})
    return {"items": result, "count": len(result)}


register(
    NodeSpec(
        type="data.select",
        group="data",
        label="Selecionar campos",
        description="Cria uma lista nova escolhendo e renomeando os campos de cada item.",
        icon="columns",
        params=(
            ParamSpec("items", "Lista", "expression", required=True, resolve=False),
            ParamSpec("mapping", "Campos", "mapping", required=True, default=[], resolve=False, item_fields=(FieldSpec("key", "Campo novo"), FieldSpec("value", "Valor", "expression", placeholder="Use ⚡ para escolher"))),
        ),
        outputs=(OutputSpec("items", "Lista", "array"), OutputSpec("count", "Quantidade", "number")),
        executor=_select,
        keywords=("map", "select", "renomear", "campos"),
    )
)


def _sort(ctx: ActionContext, params: dict[str, Any]) -> dict[str, Any]:
    items = list(_items(params.get("items")))
    key = to_text(params.get("by")).strip()

    def sort_key(item: Any) -> tuple[int, Any]:
        value = _field(item, key) if key else item
        number = to_number(value)
        return (0, number) if number is not None else (1, to_text(value).casefold())

    items.sort(key=sort_key, reverse=params.get("order") == "desc")
    limit = int(to_number(params.get("limit")) or 0)
    return {"items": items[:limit] if limit > 0 else items, "count": len(items[:limit] if limit > 0 else items)}


register(
    NodeSpec(
        type="data.sort",
        group="data",
        label="Ordenar lista",
        description="Ordena por um campo e, se quiser, fica só com os primeiros.",
        icon="sort",
        params=(
            ParamSpec("items", "Lista", "expression", required=True),
            ParamSpec("by", "Campo", "text", placeholder="valor"),
            ParamSpec("order", "Ordem", "select", default="desc", dynamic=False, options=(("asc", "Crescente"), ("desc", "Decrescente"))),
            ParamSpec("limit", "Manter só os primeiros (0 = todos)", "number", default=0, min=0),
        ),
        outputs=(OutputSpec("items", "Lista", "array"), OutputSpec("count", "Quantidade", "number")),
        executor=_sort,
        keywords=("sort", "ordenar", "top", "ranking"),
    )
)


def _group(ctx: ActionContext, params: dict[str, Any]) -> dict[str, Any]:
    items = _items(params.get("items"))
    key = to_text(params.get("by")).strip()
    sum_field = to_text(params.get("sum_field")).strip()
    if not key:
        raise ActionError("Informe o campo de agrupamento")
    groups: dict[str, list[Any]] = defaultdict(list)
    for item in items:
        groups[to_text(_field(item, key)) or "(vazio)"].append(item)
    result = []
    for name in sorted(groups):
        entry: dict[str, Any] = {"key": name, "count": len(groups[name]), "items": groups[name]}
        if sum_field:
            total = sum((Decimal(str(to_number(_field(i, sum_field)) or 0)) for i in groups[name]), Decimal(0))
            entry["total"] = float(total)
            entry["total_formatado"] = format_money(total)
        result.append(entry)
    return {"groups": result, "count": len(result)}


register(
    NodeSpec(
        type="data.group",
        group="data",
        label="Agrupar lista",
        description="Agrupa os itens por um campo (ex.: unidade), com quantidade e soma.",
        icon="layers",
        params=(
            ParamSpec("items", "Lista", "expression", required=True),
            ParamSpec("by", "Agrupar por", "text", required=True, placeholder="unidade"),
            ParamSpec("sum_field", "Somar o campo", "text", placeholder="valor"),
        ),
        outputs=(
            OutputSpec(
                "groups",
                "Grupos",
                "array",
                item_fields=(
                    OutputSpec("key", "Valor do grupo", "string"),
                    OutputSpec("count", "Quantidade", "number"),
                    OutputSpec("total", "Soma", "number"),
                    OutputSpec("total_formatado", "Soma (R$)", "string"),
                    OutputSpec("items", "Itens do grupo", "array"),
                ),
            ),
            OutputSpec("count", "Quantidade de grupos", "number"),
        ),
        executor=_group,
        keywords=("group by", "agrupar", "por unidade"),
    )
)


def _aggregate(ctx: ActionContext, params: dict[str, Any]) -> dict[str, Any]:
    items = _items(params.get("items"))
    field = to_text(params.get("field")).strip()
    operation = params.get("operation") or "sum"
    values = [to_number(_field(item, field) if field else item) for item in items]
    numbers = [Decimal(str(v)) for v in values if v is not None]
    if operation == "count":
        value: Any = len(items)
    elif not numbers:
        value = 0
    elif operation == "sum":
        value = float(sum(numbers, Decimal(0)))
    elif operation == "avg":
        value = float(sum(numbers, Decimal(0)) / len(numbers))
    elif operation == "min":
        value = float(min(numbers))
    else:
        value = float(max(numbers))
    return {"value": value, "formatted": format_money(value) if operation != "count" else str(value)}


register(
    NodeSpec(
        type="data.aggregate",
        group="data",
        label="Calcular total",
        description="Soma, média, mínimo, máximo ou contagem de um campo da lista.",
        icon="sigma",
        params=(
            ParamSpec("items", "Lista", "expression", required=True),
            ParamSpec("operation", "Cálculo", "select", default="sum", dynamic=False, options=(("sum", "Soma"), ("avg", "Média"), ("min", "Mínimo"), ("max", "Máximo"), ("count", "Contagem"))),
            ParamSpec("field", "Campo", "text", placeholder="valor"),
        ),
        outputs=(OutputSpec("value", "Resultado", "number"), OutputSpec("formatted", "Resultado (R$)", "string")),
        executor=_aggregate,
        keywords=("soma", "total", "média", "sum", "count"),
    )
)


_TABLE_HEAD = "padding:7px 9px;border-bottom:2px solid #d1d5db;font-size:12px;text-align:left;color:#6b7280"
_TABLE_CELL = "padding:7px 9px;border-bottom:1px solid #e5e7eb;font-size:13px;vertical-align:top"


def _html_table(ctx: ActionContext, params: dict[str, Any]) -> dict[str, Any]:
    items = _items(ctx.resolve(params.get("items"), None))
    columns = [entry for entry in params.get("columns") or [] if entry.get("header") or entry.get("value")]
    if not columns:
        keys: list[str] = []
        for item in items[:50]:
            for key in item if isinstance(item, dict) else []:
                if key not in keys:
                    keys.append(key)
        columns = [{"header": key, "value": f"{{{{ item.{key} }}}}"} for key in keys[:12]]
    head = "".join(f'<th style="{_TABLE_HEAD}">{escape(to_text(c.get("header")))}</th>' for c in columns)
    rows = []
    for item in items:
        cells = "".join(
            f'<td style="{_TABLE_CELL}">{escape(to_text(ctx.resolve(c.get("value"), {"item": item})))}</td>'
            for c in columns
        )
        rows.append(f"<tr>{cells}</tr>")
    html = (
        '<table role="presentation" cellspacing="0" cellpadding="0" style="border-collapse:collapse;width:100%;margin:4px 0 16px">'
        f"<thead><tr>{head}</tr></thead><tbody>{''.join(rows)}</tbody></table>"
    )
    text_rows = [" | ".join(to_text(c.get("header")) for c in columns)] + [
        " | ".join(to_text(ctx.resolve(c.get("value"), {"item": item})) for c in columns) for item in items
    ]
    csv_buffer = io.StringIO()
    writer = csv.writer(csv_buffer, delimiter=";")
    writer.writerow([to_text(c.get("header")) for c in columns])
    for item in items:
        writer.writerow([to_text(ctx.resolve(c.get("value"), {"item": item})) for c in columns])
    return {"html": html, "text": "\n".join(text_rows), "csv": csv_buffer.getvalue(), "count": len(items)}


register(
    NodeSpec(
        type="data.table",
        group="data",
        label="Criar tabela",
        description="Monta uma tabela (HTML para e-mail, texto e CSV) a partir de uma lista.",
        icon="table",
        params=(
            ParamSpec("items", "Lista", "expression", required=True, resolve=False),
            ParamSpec("columns", "Colunas (vazio = automáticas)", "columns", default=[], resolve=False, item_fields=(FieldSpec("header", "Título"), FieldSpec("value", "Valor", "expression", placeholder="Use ⚡ para escolher"))),
        ),
        outputs=(
            OutputSpec("html", "Tabela HTML", "string"),
            OutputSpec("text", "Tabela em texto", "string"),
            OutputSpec("csv", "CSV", "string"),
            OutputSpec("count", "Linhas", "number"),
        ),
        executor=_html_table,
        keywords=("html", "tabela", "csv", "relatório"),
    )
)


def _join(ctx: ActionContext, params: dict[str, Any]) -> dict[str, Any]:
    items = _items(params.get("items"))
    field = to_text(params.get("field")).strip()
    separator = to_text(params.get("separator")) if params.get("separator") not in (None, "") else ", "
    separator = separator.replace("\\n", "\n")
    values = [to_text(_field(item, field) if field else item) for item in items]
    return {"text": separator.join(value for value in values if value)}


register(
    NodeSpec(
        type="data.join",
        group="data",
        label="Juntar em texto",
        description="Junta os itens de uma lista (ou um campo deles) num texto só.",
        icon="link",
        params=(
            ParamSpec("items", "Lista", "expression", required=True),
            ParamSpec("field", "Campo (opcional)", "text", placeholder="unidade"),
            ParamSpec("separator", "Separador", "text", default=", ", placeholder=", ou \\n"),
        ),
        outputs=(OutputSpec("text", "Texto", "string"),),
        executor=_join,
        keywords=("join", "concatenar", "lista"),
    )
)
