"""Deterministic structural readers for invoice XLS and budget XLSX."""

import hashlib
import json
import re
import time
import unicodedata
from collections import Counter
from datetime import date, datetime
from datetime import time as clock_time
from decimal import Decimal
from io import BytesIO
from uuid import UUID
from zipfile import BadZipFile, ZipFile

import xlrd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from onyx.error_handling.error_codes import OnyxErrorCode
from onyx.error_handling.exceptions import OnyxError
from onyx.ton.ng_financial.models import DiagnosticLevel
from onyx.ton.ng_financial.parser import parse_date, parse_decimal
from onyx.ton.operational_import.models import (
    OperationalDiagnostic,
    OperationalKind,
    OperationalParseResult,
    OperationalRecord,
    OperationalSheetSummary,
    SheetClass,
)
from onyx.ton.sources.models import SourceFormat, SourceLocator
from onyx.utils.zip_safety import validate_zip_archive

BILLING_KEY = "billing_export"
BUDGET_ANNUAL_KEY = "budget_annual_schedule"
BUDGET_TERM_KEY = "budget_term_schedule"
PROFILE_VERSION = 1
BILLING_SOURCE_KEY = "billing_invoices"
BUDGET_SOURCE_KEY = "budget"
BILLING_COLUMNS = {
    "payer": "header:tomador",
    "invoice_number": "header:nota_fiscal",
    "emission_date": "header:emissao",
    "service_amount": "header:valor_servico",
}
BUDGET_COLUMNS = {
    "code": "A",
    "description": "B",
    "amount": "C",
    "share": "D",
}
MAX_BYTES = 10 * 1024 * 1024
MAX_SHEETS = 120
MAX_ROWS = 30_000
MAX_ROWS_PER_SHEET = 2_000
MAX_COLUMNS = 300
MAX_CELL_CHARS = 8_192
MAX_SECONDS = 45
MONTHS = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}
CODE_PATTERN = re.compile(r"^\d+(?:\.\d+)*$")


def _normal(value: object) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value).replace("\n", " "))
    return " ".join(text.encode("ascii", "ignore").decode().casefold().split())


def _raw(value: object) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _money(value: object) -> Decimal | None:
    if isinstance(value, str) and value.strip() == "-":
        return None
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        amount = Decimal(str(value))
        exponent = amount.as_tuple().exponent
        if (
            not isinstance(exponent, int)
            or exponent < -25
            or (amount != 0 and amount.adjusted() >= 25)
        ):
            raise ValueError("invalid amount")
        return amount
    return parse_decimal(value)


def _digest(parts: object) -> str:
    payload = json.dumps(
        parts, sort_keys=True, ensure_ascii=False, separators=(",", ":")
    )
    return hashlib.sha256(payload.encode()).hexdigest()


def _decimal_text(value: Decimal) -> str:
    return "0" if value == 0 else format(value.normalize(), "f")


def _diagnostic(
    result: OperationalParseResult,
    level: DiagnosticLevel,
    code: str,
    sheet: str | None = None,
    row: int | None = None,
    column: str | None = None,
) -> None:
    result.diagnostics.append(
        OperationalDiagnostic(
            level=level, code=code, sheet_name=sheet, row_number=row, column=column
        )
    )


def _check_time(start: float) -> None:
    if time.monotonic() - start > MAX_SECONDS:
        raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Operational parser time limit")


def _check_cell(value: object) -> None:
    if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
        raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Operational cell size limit")


def _empty_result(key: str, snapshot_id: UUID) -> OperationalParseResult:
    return OperationalParseResult(
        profile_key=key,
        profile_version=PROFILE_VERSION,
        snapshot_id=snapshot_id,
        records=[],
        diagnostics=[],
        sheets=[],
    )


def _field(label: str) -> str | None:
    if "tomador" in label:
        return "payer"
    if "nota fiscal" in label and ("no " in label or "n " in label):
        return "invoice_number"
    if "emissao" in label:
        return "emission_date"
    if label in ("comp", "competencia"):
        return "competence"
    if "valor" in label and "servic" in label:
        return "service_amount"
    if "deduc" in label or "reducao de base" in label:
        return "deduction_amount"
    if "liquido apos" in label or "liquido recebido" in label:
        return "net_after_discount"
    if "liquido nota" in label:
        return "invoice_net_amount"
    if "total retido" in label:
        return "total_retained"
    if "retido a mais" in label:
        return "excess_retained_amount"
    if "iss retido" in label and "proprio" not in label:
        return "iss_retained"
    if "inss retido" in label:
        return "inss_retained"
    if "ir retido" in label:
        return "ir_retained"
    if label in ("iss", "valor do iss", "iss nota fiscal"):
        return "iss_invoice"
    if label in ("inss", "valor inss", "inss nota fiscal"):
        return "inss_invoice"
    if label in ("valor ir", "valor irrf", "valor ir nota fiscal"):
        return "ir_invoice"
    if label == "periodo":
        return "period_text"
    if "tipo de recolhimento" in label or "iss retido ou proprio" in label:
        return "collection_type"
    if "status da nota" in label:
        return "invoice_status"
    return None


def _header(values: list[object]) -> dict[str, int] | None:
    mapped: dict[str, int] = {}
    for index, value in enumerate(values):
        field = _field(_normal(value))
        if field is not None and field not in mapped:
            mapped[field] = index
    return mapped if set(BILLING_COLUMNS).issubset(mapped) else None


def _invoice_number(value: object) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, float):
        if not value.is_integer():
            return None
        return str(int(value))
    text = str(value).strip()
    return text if text and text != "-" else None


def _xls_date(value: object, cell_type: int, datemode: int) -> date | None:
    if cell_type == xlrd.XL_CELL_DATE:
        if not isinstance(value, (float, int)):
            raise ValueError("invalid date")
        parsed = xlrd.xldate_as_datetime(float(value), datemode)
        if parsed.time() != clock_time(0):
            raise ValueError("invalid date")
        return parsed.date()
    return parse_date(value)


def _competence(value: object, cell_type: int, datemode: int) -> date | None:
    if value in (None, ""):
        return None
    if cell_type == xlrd.XL_CELL_DATE:
        day = _xls_date(value, cell_type, datemode)
        return date(day.year, day.month, 1) if day else None
    if isinstance(value, str):
        text = _normal(value)
        match = re.fullmatch(r"(\d{1,2})/(\d{4})", text)
        if match:
            return date(int(match.group(2)), int(match.group(1)), 1)
        match = re.fullmatch(r"([a-z]+)/([0-9]{4})", text)
        if match and match.group(1) in MONTHS:
            return date(int(match.group(2)), MONTHS[match.group(1)], 1)
        match = re.fullmatch(r"([a-z]+)/([0-9]{4}) \(\d{1,2} a \d{1,2}\)", text)
        if match and match.group(1) in MONTHS:
            return date(int(match.group(2)), MONTHS[match.group(1)], 1)
        try:
            parsed = parse_date(value)
        except ValueError:
            parsed = None
        if parsed:
            return date(parsed.year, parsed.month, 1)
    raise ValueError("invalid competence")


class BillingParser:
    key = BILLING_KEY

    def parse(  # noqa: C901 - one pass preserves the active header for each row
        self, content: bytes, snapshot_id: UUID
    ) -> OperationalParseResult:
        if not content or len(content) > MAX_BYTES:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Billing workbook size")
        start = time.monotonic()
        try:
            book = xlrd.open_workbook(file_contents=content, on_demand=True)
        except (xlrd.XLRDError, ValueError, OSError):
            raise OnyxError(
                OnyxErrorCode.INVALID_INPUT, "Invalid billing XLS"
            ) from None
        if book.nsheets > MAX_SHEETS:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Billing sheet limit")
        result = _empty_result(self.key, snapshot_id)
        duplicates: Counter[str] = Counter()
        total_rows = 0
        for sheet in book.sheets():
            _check_time(start)
            total_rows += sheet.nrows
            if (
                sheet.nrows > MAX_ROWS_PER_SHEET
                or total_rows > MAX_ROWS
                or sheet.ncols > MAX_COLUMNS
            ):
                raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Billing grid limit")
            headers = [
                (row, mapping)
                for row in range(sheet.nrows)
                if (mapping := _header(sheet.row_values(row))) is not None
            ]
            summary = OperationalSheetSummary(
                name=sheet.name,
                classification=SheetClass.DETAIL
                if headers
                else SheetClass.OUT_OF_SCOPE,
                physical_rows=sheet.nrows,
                rows_inspected=sheet.nrows,
            )
            result.sheets.append(summary)
            if not headers:
                _diagnostic(
                    result, DiagnosticLevel.INFO, "SHEET_OUT_OF_SCOPE", sheet.name
                )
                continue
            current: dict[str, int] | None = None
            for row in range(sheet.nrows):
                _check_time(start)
                values = sheet.row_values(row)
                for value in values:
                    _check_cell(value)
                mapping = _header(values)
                if mapping is not None:
                    current = mapping
                    continue
                if current is None:
                    continue
                payer = _raw(values[current["payer"]])
                invoice_value = values[current["invoice_number"]]
                date_value = values[current["emission_date"]]
                gross_value = values[current["service_amount"]]
                if not any(
                    v not in ("", None)
                    for v in (payer, invoice_value, date_value, gross_value)
                ):
                    continue
                if payer and (
                    "total" in _normal(payer)
                    or _normal(payer).startswith(("subtotal", "resumo "))
                ):
                    summary.aggregates_excluded += 1
                    continue
                if invoice_value in ("", None) and date_value in ("", None):
                    if isinstance(gross_value, str) or not payer:
                        summary.aggregates_excluded += 1
                        continue
                if (
                    invoice_value in ("", None)
                    and date_value in ("", None)
                    and gross_value in ("", None)
                ):
                    continue
                invoice = _invoice_number(invoice_value)
                if not invoice:
                    _diagnostic(
                        result,
                        DiagnosticLevel.ERROR,
                        "INVALID_INVOICE",
                        sheet.name,
                        row + 1,
                        get_column_letter(current["invoice_number"] + 1),
                    )
                    summary.rejected += 1
                    continue
                try:
                    emission = _xls_date(
                        date_value,
                        sheet.cell_type(row, current["emission_date"]),
                        book.datemode,
                    )
                except (ValueError, xlrd.XLDateError):
                    emission = None
                if emission is None:
                    _diagnostic(
                        result,
                        DiagnosticLevel.ERROR,
                        "INVALID_DATE",
                        sheet.name,
                        row + 1,
                        get_column_letter(current["emission_date"] + 1),
                    )
                    summary.rejected += 1
                    continue
                try:
                    gross = _money(gross_value)
                except ValueError:
                    gross = None
                if gross is None:
                    _diagnostic(
                        result,
                        DiagnosticLevel.ERROR,
                        "INVALID_AMOUNT",
                        sheet.name,
                        row + 1,
                        get_column_letter(current["service_amount"] + 1),
                    )
                    summary.rejected += 1
                    continue
                competence: date | None = None
                if "competence" in current:
                    index = current["competence"]
                    try:
                        competence = _competence(
                            values[index], sheet.cell_type(row, index), book.datemode
                        )
                    except (ValueError, xlrd.XLDateError):
                        _diagnostic(
                            result,
                            DiagnosticLevel.ERROR,
                            "INVALID_COMPETENCE",
                            sheet.name,
                            row + 1,
                            get_column_letter(index + 1),
                        )
                        summary.rejected += 1
                        continue
                typed: dict[str, str | None] = {}
                numeric: dict[str, Decimal | None] = {}
                failed_column: str | None = None
                for field, index in current.items():
                    if field in (
                        "payer",
                        "invoice_number",
                        "emission_date",
                        "competence",
                        "service_amount",
                    ):
                        continue
                    value = values[index]
                    if field.endswith("_amount") or field in (
                        "total_retained",
                        "iss_retained",
                        "inss_retained",
                        "ir_retained",
                        "iss_invoice",
                        "inss_invoice",
                        "ir_invoice",
                    ):
                        try:
                            amount = _money(value)
                            numeric[field] = amount
                        except ValueError:
                            failed_column = get_column_letter(index + 1)
                            break
                    else:
                        typed[field] = _raw(value)
                if failed_column:
                    _diagnostic(
                        result,
                        DiagnosticLevel.ERROR,
                        "INVALID_AMOUNT",
                        sheet.name,
                        row + 1,
                        failed_column,
                    )
                    summary.rejected += 1
                    continue
                if not payer:
                    _diagnostic(
                        result,
                        DiagnosticLevel.WARNING,
                        "PAYER_BLANK",
                        sheet.name,
                        row + 1,
                        get_column_letter(current["payer"] + 1),
                    )
                if competence is not None and abs(competence.year - emission.year) > 5:
                    _diagnostic(
                        result,
                        DiagnosticLevel.WARNING,
                        "COMPETENCE_YEAR_GAP",
                        sheet.name,
                        row + 1,
                        get_column_letter(current["competence"] + 1),
                    )
                retention_fields = (
                    "iss_retained",
                    "inss_retained",
                    "ir_retained",
                )
                retention_values = [numeric.get(field) for field in retention_fields]
                if (
                    all(value is not None for value in retention_values)
                    and numeric.get("total_retained") is not None
                ):
                    retention_sum = sum(
                        (value for value in retention_values if value is not None),
                        Decimal(0),
                    )
                    total_retained = numeric["total_retained"]
                    assert total_retained is not None
                    if abs(retention_sum - total_retained) > Decimal("0.02"):
                        _diagnostic(
                            result,
                            DiagnosticLevel.WARNING,
                            "RETENTION_TOTAL_MISMATCH",
                            sheet.name,
                            row + 1,
                            get_column_letter(current["total_retained"] + 1),
                        )
                source_values = {
                    get_column_letter(index + 1): _raw(value)
                    for index, value in enumerate(values)
                }
                fingerprint = _digest(
                    [
                        self.key,
                        PROFILE_VERSION,
                        payer,
                        invoice,
                        emission.isoformat(),
                        competence.isoformat() if competence else None,
                        _decimal_text(gross),
                        typed,
                        {
                            field: _decimal_text(value) if value is not None else None
                            for field, value in numeric.items()
                        },
                    ]
                )
                duplicates[fingerprint] += 1
                result.records.append(
                    OperationalRecord(
                        locator=SourceLocator(
                            snapshot_id=snapshot_id,
                            sheet_name=sheet.name,
                            row_number=row + 1,
                        ),
                        kind=OperationalKind.BILLING,
                        identifier=invoice,
                        description=payer,
                        record_date=emission,
                        competence=competence,
                        amount=gross,
                        period_basis=None,
                        numeric_values=numeric,
                        typed_values=typed,
                        source_values=source_values,
                        fingerprint=fingerprint,
                        duplicate_ordinal=duplicates[fingerprint],
                    )
                )
                summary.accepted += 1
        if not result.records:
            raise OnyxError(
                OnyxErrorCode.INVALID_INPUT, "Unrecognized billing workbook"
            )
        return result


def _budget_key(sheet: Worksheet) -> str | None:
    annual = _normal(sheet.cell(2, 1).value)
    monthly = _normal(sheet.cell(3, 1).value)
    if "valor do contrato" not in annual or "mensal" not in monthly:
        return None
    if (
        _normal(sheet.cell(8, 1).value) != "item"
        or _normal(sheet.cell(8, 2).value) != "descricao"
        or _normal(sheet.cell(8, 3).value) != "valor total"
    ):
        return None
    if "composicao de custos" not in _normal(sheet.cell(10, 2).value):
        return None
    if "anual" in annual and sheet.cell(3, 3).value == "=C2/12":
        return BUDGET_ANNUAL_KEY
    if "meses" in annual and sheet.cell(2, 3).value == "=C3*30":
        return BUDGET_TERM_KEY
    return None


def _sheet_class(name: str) -> SheetClass:
    normalized = _normal(name)
    if any(
        word in normalized
        for word in ("sintese", "resumo", "cronograma", "planilha de precos", "ppu")
    ):
        return SheetClass.SUMMARY
    if any(
        word in normalized for word in ("preco", "bdi", "salario", "cotacao", "banco")
    ):
        return SheetClass.REFERENCE
    return SheetClass.SUPPORT


class BudgetParser:
    def __init__(self, key: str):
        if key not in (BUDGET_ANNUAL_KEY, BUDGET_TERM_KEY):
            raise ValueError("Unknown budget profile")
        self.key = key

    def parse(self, content: bytes, snapshot_id: UUID) -> OperationalParseResult:
        if not content or len(content) > MAX_BYTES:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Budget workbook size")
        start = time.monotonic()
        try:
            with ZipFile(BytesIO(content)) as archive:
                validate_zip_archive(archive)
            formulas = load_workbook(
                BytesIO(content), read_only=False, data_only=False, keep_links=False
            )
            cached = load_workbook(
                BytesIO(content), read_only=False, data_only=True, keep_links=False
            )
        except (BadZipFile, OSError, ValueError, KeyError):
            raise OnyxError(
                OnyxErrorCode.INVALID_INPUT, "Invalid budget XLSX"
            ) from None
        if len(formulas.worksheets) > MAX_SHEETS:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Budget sheet limit")
        candidates = [
            sheet for sheet in formulas.worksheets if _budget_key(sheet) is not None
        ]
        if len(candidates) != 1 or _budget_key(candidates[0]) != self.key:
            raise OnyxError(
                OnyxErrorCode.INVALID_INPUT, "Unrecognized budget structure"
            )
        detail = candidates[0]
        result = _empty_result(self.key, snapshot_id)
        for sheet in formulas.worksheets:
            _check_time(start)
            result.sheets.append(
                OperationalSheetSummary(
                    name=sheet.title,
                    classification=SheetClass.DETAIL
                    if sheet is detail
                    else _sheet_class(sheet.title),
                    physical_rows=sheet.max_row,
                    rows_inspected=sheet.max_row if sheet is detail else 0,
                )
            )
        summary = next(item for item in result.sheets if item.name == detail.title)
        if detail.max_row > MAX_ROWS_PER_SHEET or detail.max_column > MAX_COLUMNS:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Budget grid limit")
        finish = next(
            (
                row
                for row in range(11, detail.max_row + 1)
                if "composicao da proposta" in _normal(detail.cell(row, 2).value)
            ),
            None,
        )
        if finish is None:
            raise OnyxError(
                OnyxErrorCode.INVALID_INPUT, "Budget section boundary absent"
            )
        summary.proposal_rows_excluded = sum(
            CODE_PATTERN.fullmatch(str(detail.cell(row, 1).value).strip()) is not None
            for row in range(finish, detail.max_row + 1)
        )
        codes = {
            row: str(detail.cell(row, 1).value).strip()
            for row in range(10, finish)
            if CODE_PATTERN.fullmatch(str(detail.cell(row, 1).value).strip())
        }
        duplicates: Counter[str] = Counter()
        for row, code in codes.items():
            _check_time(start)
            if any(
                other != row and value.startswith(code + ".")
                for other, value in codes.items()
            ):
                summary.aggregates_excluded += 1
                continue
            description = detail.cell(row, 2).value
            raw_amount = detail.cell(row, 3).value
            cached_amount = cached[detail.title].cell(row, 3).value
            if any(
                isinstance(detail.cell(row, column), MergedCell)
                for column in (1, 2, 3, 4)
            ):
                _diagnostic(
                    result,
                    DiagnosticLevel.ERROR,
                    "MERGED_DETAIL_UNSUPPORTED",
                    detail.title,
                    row,
                    "A:C",
                )
                summary.rejected += 1
                continue
            if raw_amount is None:
                _diagnostic(
                    result,
                    DiagnosticLevel.ERROR,
                    "MISSING_AMOUNT",
                    detail.title,
                    row,
                    "C",
                )
                summary.rejected += 1
                continue
            if (
                isinstance(raw_amount, str)
                and raw_amount.startswith("=")
                and cached_amount is None
            ):
                _diagnostic(
                    result,
                    DiagnosticLevel.ERROR,
                    "FORMULA_CACHE_MISSING",
                    detail.title,
                    row,
                    "C",
                )
                summary.rejected += 1
                continue
            try:
                amount = _money(cached_amount)
            except ValueError:
                amount = None
            if amount is None:
                _diagnostic(
                    result,
                    DiagnosticLevel.ERROR,
                    "INVALID_AMOUNT",
                    detail.title,
                    row,
                    "C",
                )
                summary.rejected += 1
                continue
            source_values = {
                get_column_letter(column): _raw(detail.cell(row, column).value)
                for column in range(1, 5)
            }
            share_value = cached[detail.title].cell(row, 4).value
            try:
                share = _money(share_value)
            except ValueError:
                share = None
                _diagnostic(
                    result,
                    DiagnosticLevel.WARNING,
                    "INVALID_SHARE",
                    detail.title,
                    row,
                    "D",
                )
            fingerprint = _digest(
                [
                    self.key,
                    PROFILE_VERSION,
                    code,
                    _raw(description),
                    _decimal_text(amount),
                ]
            )
            duplicates[fingerprint] += 1
            result.records.append(
                OperationalRecord(
                    locator=SourceLocator(
                        snapshot_id=snapshot_id, sheet_name=detail.title, row_number=row
                    ),
                    kind=OperationalKind.BUDGET,
                    identifier=code,
                    description=_raw(description),
                    record_date=None,
                    competence=None,
                    amount=amount,
                    period_basis="MONTHLY_CONTRACT",
                    numeric_values={"share": share},
                    typed_values={
                        "contract_label": _raw(detail.cell(1, 1).value),
                        "contract_term_months": "12"
                        if self.key == BUDGET_ANNUAL_KEY
                        else "30",
                    },
                    source_values=source_values,
                    formula_cached=isinstance(raw_amount, str)
                    and raw_amount.startswith("="),
                    fingerprint=fingerprint,
                    duplicate_ordinal=duplicates[fingerprint],
                )
            )
            summary.accepted += 1
        if not result.records:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Budget detail absent")
        return result


def identify_profile(content: bytes, format: SourceFormat) -> str:
    """Select by workbook structure. Names and acquisition method are ignored."""
    if format == SourceFormat.XLS:
        return BillingParser().parse(content, UUID(int=0)).profile_key
    if format == SourceFormat.XLSX:
        if not content or len(content) > MAX_BYTES:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Budget workbook size")
        try:
            with ZipFile(BytesIO(content)) as archive:
                validate_zip_archive(archive)
            book = load_workbook(
                BytesIO(content), read_only=False, data_only=False, keep_links=False
            )
        except (BadZipFile, OSError, ValueError, KeyError):
            raise OnyxError(
                OnyxErrorCode.INVALID_INPUT, "Invalid budget XLSX"
            ) from None
        if len(book.worksheets) > MAX_SHEETS:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Budget sheet limit")
        keys = [key for sheet in book.worksheets if (key := _budget_key(sheet))]
        if len(keys) == 1:
            return keys[0]
    raise OnyxError(OnyxErrorCode.INVALID_INPUT, "No operational profile matches")
