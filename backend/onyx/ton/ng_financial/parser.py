"""NG financial worksheet v1. No business value is corrected here."""

import hashlib
import json
import re
import time
from collections import Counter
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from datetime import time as clock_time
from decimal import Decimal, InvalidOperation
from io import BytesIO
from uuid import UUID
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from onyx.error_handling.error_codes import OnyxErrorCode
from onyx.error_handling.exceptions import OnyxError
from onyx.ton.ng_financial.models import (
    DiagnosticCode,
    DiagnosticLevel,
    ParseDiagnostic,
    ParsedImportResult,
    ParsedSourceRecordData,
    RowKind,
    SheetSummary,
)
from onyx.ton.sources.models import SourceLocator
from onyx.utils.zip_safety import validate_zip_archive

PROFILE_KEY = "ng_financial_export"
PROFILE_VERSION = 1
SOURCE_KEY = "financial_launches"
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_SHEETS = 24
MAX_ROWS_PER_SHEET = 10_000
MAX_TOTAL_ROWS = 50_000
MAX_COLUMNS = 32
MAX_CELL_CHARS = 8192
MAX_SECONDS = 30
AMOUNT_PRECISION = 30
AMOUNT_SCALE = 10
MONTHS = {
    "jan": 1,
    "fev": 2,
    "mar": 3,
    "abr": 4,
    "mai": 5,
    "jun": 6,
    "jul": 7,
    "ago": 8,
    "set": 9,
    "out": 10,
    "nov": 11,
    "dez": 12,
}
SHEET_PATTERN = re.compile(
    r"^(jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)(?:[ ._-]+ok)?$", re.I
)
ACCOUNT_PATTERN = re.compile(r"^\s*(\d+(?:\.\d+)*)\s*-\s*(\S.*)$", re.S)
# pt-BR only: dot groups thousands, comma separates decimals.
BRAZILIAN_NUMBER = re.compile(r"^-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?$")
AMOUNT_FIELDS = (
    "movement_amount",
    "movement_retention_amount",
    "movement_net_amount",
    "installment_retention_amount",
    "installment_net_amount",
    "interest_amount",
    "penalty_amount",
    "discount_amount",
    "expense_amount",
    "loss_amount",
    "other_deduction_amount",
    "final_amount",
)
CONTRACT_WIDTH = 17
COLUMNS = tuple(get_column_letter(index) for index in range(1, CONTRACT_WIDTH + 1))
COLUMN_MAP: dict[str, str] = dict(
    zip(
        COLUMNS,
        (
            "account_path",
            "emission_date",
            "administrative_unit",
            "document_number",
            "history",
            *AMOUNT_FIELDS,
        ),
        strict=True,
    )
)
# Labels from the consolidated reference workbook. The NG export itself has no
# header row; these are only used to recognise one if a later export adds it.
HEADER = (
    "Código Estruturado",
    "Emissão",
    "Unidade Administrativa",
    "Número",
    "Histórico",
    "Valor Movimento",
    "Valor Imposto Retido Movimento",
    "Valor Líquido Movimento",
    "Valor Imposto Retido Parcelas",
    "Valor Líquido Parcelas",
    "Valor Juros",
    "Valor Multa",
    "Valor Desconto",
    "Valor Despesa",
    "Valor Perda",
    "Valor Outras Deduções",
    "Valor Final",
)
AMOUNT_SLICE = slice(5, CONTRACT_WIDTH)


def _bounded_amount(result: Decimal) -> Decimal:
    # Must fit TON_AMOUNT (30, 10) exactly; the database would otherwise round.
    exponent = result.as_tuple().exponent
    if (
        not isinstance(exponent, int)
        or exponent < -AMOUNT_SCALE
        or (result != 0 and result.adjusted() >= AMOUNT_PRECISION - AMOUNT_SCALE)
    ):
        raise ValueError("invalid amount")
    return result


def parse_decimal(value: object) -> Decimal | None:
    """Exact amount. Blank is None; zero stays zero; ambiguity is an error."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("invalid amount")
    if isinstance(value, (int, float, Decimal)):
        try:
            return _bounded_amount(Decimal(str(value)))
        except InvalidOperation:
            raise ValueError("invalid amount") from None
    if not isinstance(value, str):
        raise ValueError("invalid amount")
    text = value.replace("\u00a0", " ").strip()
    if text == "":
        return None
    if text.startswith("R$"):
        text = text[2:].strip()
    if not BRAZILIAN_NUMBER.fullmatch(text):
        raise ValueError("invalid amount")
    try:
        return _bounded_amount(Decimal(text.replace(".", "").replace(",", ".")))
    except InvalidOperation:
        raise ValueError("invalid amount") from None


def parse_date(value: object) -> date | None:
    """Excel date cells or explicit day-first / ISO text. No serial guessing."""
    if value is None:
        return None
    if isinstance(value, datetime):
        if value.time() != clock_time(0):
            raise ValueError("invalid date")
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        if text == "":
            return None
        for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, pattern).date()
            except ValueError:
                pass
    raise ValueError("invalid date")


def sheet_month(name: str) -> int | None:
    match = SHEET_PATTERN.fullmatch(name.strip())
    return MONTHS[match.group(1).lower()] if match else None


def _as_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, (date, clock_time)):
        return value.isoformat()
    return str(value)


def _canonical_decimal(value: Decimal | None) -> str | None:
    if value is None:
        return None
    if value == 0:
        return "0"
    return format(value.normalize(), "f")


def _comparison_value(value: object) -> str | None:
    """Type-tagged cell identity used only for hierarchy reconciliation."""
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return "n:" + (_canonical_decimal(Decimal(str(value))) or "")
    return f"{type(value).__name__}:{value}"


def _account(value: object) -> tuple[str, str] | None:
    if not isinstance(value, str):
        return None
    match = ACCOUNT_PATTERN.fullmatch(value)
    return (match.group(1), match.group(2).strip()) if match else None


def _direct_parent(code: str, codes: set[str]) -> str | None:
    parts = code.split(".")
    for size in range(len(parts) - 1, 0, -1):
        candidate = ".".join(parts[:size])
        if candidate in codes:
            return candidate
    return None


@dataclass(frozen=True)
class _SheetRow:
    number: int
    values: tuple[object, ...]
    extras: tuple[tuple[str, object], ...]
    formula_column: str | None

    @property
    def is_blank(self) -> bool:
        return not self.extras and all(value is None for value in self.values)


class _Budget:
    def __init__(self) -> None:
        self.deadline = time.monotonic() + MAX_SECONDS
        self.rows = 0

    def tick(self) -> None:
        if time.monotonic() > self.deadline:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "NG parser time limit")


def _bounded_workbook(content: bytes) -> None:
    if not content or len(content) > MAX_WORKBOOK_BYTES:
        raise OnyxError(OnyxErrorCode.INVALID_INPUT, "NG workbook size is invalid")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = validate_zip_archive(archive)
            names = {entry.filename.lower() for entry in entries}
    except BadZipFile:
        raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Invalid NG workbook") from None
    if "xl/workbook.xml" not in names or any("vbaproject" in name for name in names):
        raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Invalid NG workbook")


def _read_sheet(sheet: ReadOnlyWorksheet, budget: _Budget) -> list[_SheetRow]:
    """Stream every stored row. Declared dimensions are not trusted."""
    sheet.reset_dimensions()
    rows: list[_SheetRow] = []
    for number, cells in enumerate(sheet.iter_rows(), start=1):
        budget.tick()
        budget.rows += 1
        if number > MAX_ROWS_PER_SHEET or budget.rows > MAX_TOTAL_ROWS:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "NG workbook row limit")
        values: list[object] = [None] * CONTRACT_WIDTH
        extras: list[tuple[str, object]] = []
        formula_column: str | None = None
        for position, cell in enumerate(cells):
            if isinstance(cell, EmptyCell):
                continue
            assert isinstance(cell, ReadOnlyCell)
            value = cell.value
            if value is None:
                continue
            if position >= MAX_COLUMNS:
                raise OnyxError(
                    OnyxErrorCode.INVALID_INPUT, "NG worksheet column limit"
                )
            if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
                raise OnyxError(OnyxErrorCode.INVALID_INPUT, "NG cell size limit")
            if cell.data_type == "f" and formula_column is None:
                formula_column = get_column_letter(position + 1)
            if position < CONTRACT_WIDTH:
                values[position] = value
            else:
                extras.append((get_column_letter(position + 1), value))
        rows.append(_SheetRow(number, tuple(values), tuple(extras), formula_column))
    return rows


def _is_header(values: Sequence[object]) -> bool:
    return all(
        isinstance(value, str) and value.strip().casefold() == label.casefold()
        for value, label in zip(values, HEADER, strict=True)
    )


def _is_launch_shape(values: Sequence[object]) -> bool:
    date_value = values[1]
    try:
        has_date = parse_date(date_value) is not None
    except ValueError:
        has_date = False
    return (
        has_date
        and values[4] is not None
        and any(
            isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)
            for value in values[AMOUNT_SLICE]
        )
    )


def _recognised(rows: Sequence[_SheetRow]) -> bool:
    """Structural contract: account-led sections with launch-shaped rows."""
    content = [row for row in rows if not row.is_blank]
    first = content[0].values
    if _account(first[0]) is None and not _is_header(first):
        return False
    codes = {
        account[0]
        for row in content
        if (account := _account(row.values[0])) is not None
    }
    return any("." in code for code in codes) and any(
        _account(row.values[0]) is not None and _is_launch_shape(row.values)
        for row in content
    )


def _classify(
    row: _SheetRow, current_code: str | None, parents: set[str], header_seen: bool
) -> RowKind:
    values = row.values
    if row.is_blank:
        return RowKind.BLANK
    if _is_header(values):
        return RowKind.REPEATED_HEADER if header_seen else RowKind.HEADER
    first = values[0].strip().casefold() if isinstance(values[0], str) else ""
    history = values[4].strip().casefold() if isinstance(values[4], str) else ""
    if first.startswith("total geral") or first == "total":
        return RowKind.TOTAL
    if first.startswith("subtotal") or history.startswith("subtotal"):
        return RowKind.SUBTOTAL
    account = _account(values[0])
    if values[0] is not None and account is None:
        return RowKind.UNKNOWN
    if current_code is None:
        return RowKind.UNKNOWN
    if values[4] is None:
        if account is not None and all(value is None for value in values[1:]):
            return RowKind.SECTION
        return RowKind.UNKNOWN
    if current_code in parents:
        return RowKind.HIERARCHY
    return RowKind.DETAIL


def _fingerprint(
    account_code: str,
    account_label: str,
    emission_date: date,
    unit: str | None,
    document: str | None,
    history: str,
    amounts: dict[str, Decimal | None],
    extras: Iterable[tuple[str, str | None]],
) -> str:
    payload = json.dumps(
        [
            PROFILE_KEY,
            PROFILE_VERSION,
            account_code,
            account_label,
            emission_date.isoformat(),
            unit,
            document,
            history,
            [_canonical_decimal(amounts[name]) for name in AMOUNT_FIELDS],
            sorted(extras),
        ],
        ensure_ascii=False,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode()).hexdigest()


@dataclass
class _GroupState:
    """Carry-forward state inside one account block.

    The export groups rows: B starts a date group and C starts a unit group
    nested inside it. A blank B continues the date. A new date with blank C
    has no unit; C is never carried across dates or accounts.
    """

    code: str | None = None
    label: str = ""
    emission_date: date | None = None
    date_invalid: bool = False
    unit: str | None = None

    def advance(self, values: Sequence[object]) -> DiagnosticCode | None:
        error: DiagnosticCode | None = None
        if values[1] is not None:
            self.unit = None
            try:
                self.emission_date = parse_date(values[1])
                self.date_invalid = False
                if self.emission_date is None:
                    error = DiagnosticCode.MISSING_DATE
            except ValueError:
                # Continuation rows must not inherit an earlier valid date.
                self.emission_date, self.date_invalid = None, True
                error = DiagnosticCode.INVALID_DATE
        elif self.date_invalid:
            error = DiagnosticCode.INHERITED_DATE_INVALID
        elif self.emission_date is None:
            error = DiagnosticCode.MISSING_DATE
        if values[2] is not None:
            self.unit = str(values[2])
        return error


class _SheetParser:
    """One accepted monthly worksheet."""

    def __init__(
        self,
        name: str,
        month: int,
        rows: list[_SheetRow],
        snapshot_id: UUID,
        result: ParsedImportResult,
    ) -> None:
        self.name = name
        self.month = month
        self.rows = rows
        self.snapshot_id = snapshot_id
        self.result = result
        self.duplicates: Counter[str] = Counter()

    def _diagnose(
        self,
        level: DiagnosticLevel,
        code: DiagnosticCode,
        row_number: int | None = None,
        column: str | None = None,
    ) -> None:
        self.result.diagnostics.append(
            ParseDiagnostic(
                level=level,
                code=code,
                sheet_name=self.name,
                row_number=row_number,
                column=column,
            )
        )

    def _reject(
        self, code: DiagnosticCode, row_number: int, column: str | None = None
    ) -> None:
        self._diagnose(DiagnosticLevel.ERROR, code, row_number, column)

    def _collect_codes(self) -> set[str]:
        codes: set[str] = set()
        for row in self.rows:
            account = _account(row.values[0])
            if account is None:
                continue
            if account[0] in codes:
                self._diagnose(
                    DiagnosticLevel.WARNING,
                    DiagnosticCode.ACCOUNT_SECTION_REPEATED,
                    row.number,
                    "A",
                )
            codes.add(account[0])
        return codes

    def parse(self, budget: _Budget) -> SheetSummary:
        codes = self._collect_codes()
        # A code with any child section is a parent; its rows repeat children.
        parents = {
            parent
            for code in codes
            if (parent := _direct_parent(code, codes)) is not None
        }
        counts: Counter[RowKind] = Counter()
        spans: dict[RowKind, list[tuple[int, int]]] = {}
        block_rows: dict[str, list[tuple[int, tuple[str | None, ...]]]] = {}
        state = _GroupState()
        header_seen = False
        for row in self.rows:
            budget.tick()
            account = _account(row.values[0])
            if account is not None:
                state = _GroupState(code=account[0], label=account[1])
            kind = _classify(row, state.code, parents, header_seen)
            header_seen = header_seen or kind in (
                RowKind.HEADER,
                RowKind.REPEATED_HEADER,
            )
            counts[kind] += 1
            ranges = spans.setdefault(kind, [])
            if ranges and ranges[-1][1] == row.number - 1:
                ranges[-1] = (ranges[-1][0], row.number)
            else:
                ranges.append((row.number, row.number))
            for column, _value in row.extras:
                self._diagnose(
                    DiagnosticLevel.INFO,
                    DiagnosticCode.CELL_OUTSIDE_CONTRACT,
                    row.number,
                    column,
                )
            if kind == RowKind.UNKNOWN:
                self._reject(DiagnosticCode.UNKNOWN_ROW, row.number)
            elif kind in (RowKind.DETAIL, RowKind.HIERARCHY):
                assert state.code is not None
                date_error = state.advance(row.values)
                block_rows.setdefault(state.code, []).append(
                    (
                        row.number,
                        tuple(
                            _comparison_value(value)
                            for value in row.values[AMOUNT_SLICE]
                        ),
                    )
                )
                if kind == RowKind.DETAIL:
                    self._detail(row, state, date_error)
                elif row.formula_column is not None:
                    self._diagnose(
                        DiagnosticLevel.WARNING,
                        DiagnosticCode.FORMULA_UNSUPPORTED,
                        row.number,
                        row.formula_column,
                    )
        self._reconcile(codes, block_rows)
        return SheetSummary(
            name=self.name,
            month=self.month,
            accepted=True,
            physical_rows=len(self.rows),
            classifications=dict(counts),
            row_spans=spans,
        )

    def _amounts(self, row: _SheetRow) -> dict[str, Decimal | None] | None:
        amounts: dict[str, Decimal | None] = {}
        for offset, name in enumerate(AMOUNT_FIELDS):
            try:
                amounts[name] = parse_decimal(row.values[offset + 5])
            except ValueError:
                self._reject(
                    DiagnosticCode.INVALID_AMOUNT, row.number, COLUMNS[offset + 5]
                )
                return None
        if all(amount is None for amount in amounts.values()):
            self._reject(DiagnosticCode.MISSING_AMOUNTS, row.number)
            return None
        return amounts

    def _detail(
        self, row: _SheetRow, state: _GroupState, date_error: DiagnosticCode | None
    ) -> None:
        if row.formula_column is not None:
            self._reject(
                DiagnosticCode.FORMULA_UNSUPPORTED, row.number, row.formula_column
            )
            return
        if date_error is not None:
            self._reject(date_error, row.number, "B")
            return
        amounts = self._amounts(row)
        if amounts is None:
            return
        assert state.code is not None and state.emission_date is not None
        if state.unit is None:
            self._diagnose(
                DiagnosticLevel.WARNING, DiagnosticCode.UNIT_BLANK, row.number, "C"
            )
        values = row.values
        source_values = {
            column: _as_text(values[index]) for index, column in enumerate(COLUMNS)
        }
        extras = [(column, _as_text(value)) for column, value in row.extras]
        source_values.update(extras)
        document = _as_text(values[3])
        history = str(values[4])
        fingerprint = _fingerprint(
            state.code,
            state.label,
            state.emission_date,
            state.unit,
            document,
            history,
            amounts,
            extras,
        )
        self.duplicates[fingerprint] += 1
        self.result.records.append(
            ParsedSourceRecordData(
                locator=SourceLocator(
                    snapshot_id=self.snapshot_id,
                    sheet_name=self.name,
                    row_number=row.number,
                ),
                sheet_month=self.month,
                account_code=state.code,
                account_label=state.label,
                emission_date=state.emission_date,
                administrative_unit=state.unit,
                document_number=document,
                history=history,
                source_values=source_values,
                fingerprint=fingerprint,
                duplicate_ordinal=self.duplicates[fingerprint],
                **amounts,
            )
        )

    def _reconcile(
        self,
        codes: set[str],
        block_rows: dict[str, list[tuple[int, tuple[str | None, ...]]]],
    ) -> None:
        """Compare each parent block with its direct children on F:Q.

        Parent rows repeat child launches, so only leaf rows are records.
        Unmatched rows are reported by location and never imported or fixed.
        """
        children: dict[str, list[str]] = {}
        for code in sorted(codes):
            parent = _direct_parent(code, codes)
            if parent is not None:
                children.setdefault(parent, []).append(code)
        for parent in sorted(children):
            parent_rows = block_rows.get(parent, [])
            child_rows = sorted(
                (
                    item
                    for child in children[parent]
                    for item in block_rows.get(child, [])
                ),
                key=lambda item: item[0],
            )
            available = Counter(key for _number, key in child_rows)
            for number, key in parent_rows:
                if available[key] > 0:
                    available[key] -= 1
                else:
                    self._diagnose(
                        DiagnosticLevel.WARNING,
                        DiagnosticCode.PARENT_ROW_WITHOUT_CHILD_MATCH,
                        number,
                    )
            remaining = Counter(key for _number, key in parent_rows)
            for number, key in child_rows:
                if remaining[key] > 0:
                    remaining[key] -= 1
                else:
                    self._diagnose(
                        DiagnosticLevel.WARNING,
                        DiagnosticCode.CHILD_ROW_WITHOUT_PARENT_MATCH,
                        number,
                    )


class NgFinancialExportParser:
    def parse(self, content: bytes, snapshot_id: UUID) -> ParsedImportResult:
        _bounded_workbook(content)
        budget = _Budget()
        try:
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=False,
                keep_links=False,
                keep_vba=False,
            )
        except Exception:
            raise OnyxError(
                OnyxErrorCode.INVALID_INPUT, "Invalid NG workbook"
            ) from None
        try:
            if len(workbook.sheetnames) > MAX_SHEETS:
                raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Too many workbook sheets")
            result = ParsedImportResult(
                profile_key=PROFILE_KEY,
                profile_version=PROFILE_VERSION,
                snapshot_id=snapshot_id,
                records=[],
                diagnostics=[],
                sheets=[],
            )
            months_seen: set[int] = set()
            for sheet in workbook.worksheets:
                budget.tick()
                assert isinstance(sheet, ReadOnlyWorksheet)
                month = sheet_month(sheet.title)
                if month is None:
                    result.sheets.append(_skipped(sheet.title, None))
                    result.diagnostics.append(
                        ParseDiagnostic(
                            level=DiagnosticLevel.INFO,
                            code=DiagnosticCode.SHEET_OUT_OF_SCOPE,
                            sheet_name=sheet.title,
                        )
                    )
                    continue
                rows = _read_sheet(sheet, budget)
                if all(row.is_blank for row in rows):
                    result.sheets.append(_skipped(sheet.title, month, len(rows)))
                    result.diagnostics.append(
                        ParseDiagnostic(
                            level=DiagnosticLevel.WARNING,
                            code=DiagnosticCode.MONTH_SHEET_EMPTY,
                            sheet_name=sheet.title,
                        )
                    )
                    continue
                if not _recognised(rows):
                    raise OnyxError(
                        OnyxErrorCode.INVALID_INPUT,
                        "Workbook does not match NG financial profile",
                    )
                if month in months_seen:
                    raise OnyxError(
                        OnyxErrorCode.INVALID_INPUT,
                        "Workbook has two worksheets for one month",
                    )
                months_seen.add(month)
                summary = _SheetParser(
                    sheet.title, month, rows, snapshot_id, result
                ).parse(budget)
                result.sheets.append(summary)
            if not months_seen:
                raise OnyxError(
                    OnyxErrorCode.INVALID_INPUT,
                    "Workbook does not match NG financial profile",
                )
            return result
        finally:
            workbook.close()


def _skipped(name: str, month: int | None, physical_rows: int = 0) -> SheetSummary:
    return SheetSummary(
        name=name,
        month=month,
        accepted=False,
        physical_rows=physical_rows,
        classifications={},
        row_spans={},
    )
