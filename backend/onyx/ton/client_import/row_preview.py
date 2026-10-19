"""The cells of a few spreadsheet rows, read back from the imported file.

The stored diagnostics only say where a problem is (sheet, row, column). To
fix it the Controladoria needs to see the row itself, so this reads those rows
from the same file the import used, with the header above each one.
"""

import datetime
from collections import defaultdict
from io import BytesIO

import xlrd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from onyx.ton.sources.models import SourceFormat

MAX_CELLS = 24
HEADER_MIN_TEXT_CELLS = 4


class RowCell(BaseModel):
    column: str
    header: str | None
    value: str


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0):
            return value.strftime("%d/%m/%Y")
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, datetime.date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.2f}".replace(".", ",")
    return str(value).strip()


def _is_header(values: list[object]) -> bool:
    filled = [value for value in values if value not in (None, "")]
    texts = [
        value
        for value in filled
        if isinstance(value, str) and not value.strip().replace(".", "").isdigit()
    ]
    return len(texts) >= HEADER_MIN_TEXT_CELLS and len(texts) >= 0.7 * len(filled)


def _cells(values: list[object], header: list[object] | None) -> list[RowCell]:
    cells: list[RowCell] = []
    for index, value in enumerate(values):
        text = _text(value)
        if not text:
            continue
        title = _text(header[index]) if header and index < len(header) else ""
        cells.append(
            RowCell(
                column=get_column_letter(index + 1), header=title or None, value=text
            )
        )
        if len(cells) >= MAX_CELLS:
            break
    return cells


def _scan(
    rows: list[tuple[int, list[object]]],
    wanted: set[int],
    default_header: tuple[str, ...] | None = None,
) -> dict[int, list[RowCell]]:
    """Cells of the wanted rows, with the header in force above each one.

    Without a header row (the NG export), ``default_header`` names the columns
    and the account of the block, written only on its first row, is added to
    rows that leave column A empty.
    """
    found: dict[int, list[RowCell]] = {}
    header: list[object] | None = list(default_header) if default_header else None
    section: str | None = None
    for number, values in rows:
        first = _text(values[0]) if values else ""
        if default_header and first:
            section = first
        if number in wanted:
            cells = _cells(values, header)
            if default_header and section and not first:
                cells.insert(
                    0,
                    RowCell(column="A", header=f"{default_header[0]} (bloco)", value=section),
                )
            found[number] = cells
        if _is_header(values):
            header = values
        if len(found) == len(wanted):
            break
    return found


def read_rows(
    content: bytes,
    file_format: SourceFormat,
    wanted: dict[str, set[int]],
    default_header: tuple[str, ...] | None = None,
) -> dict[tuple[str, int], list[RowCell]]:
    """Cells of the wanted (sheet, 1-based row) pairs; missing ones are left out."""
    result: dict[tuple[str, int], list[RowCell]] = {}
    if not wanted:
        return result
    if file_format == SourceFormat.XLS:
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        try:
            for name, numbers in wanted.items():
                if name not in book.sheet_names():
                    continue
                sheet = book.sheet_by_name(name)
                last = min(max(numbers), sheet.nrows)

                def values(row: int) -> list[object]:
                    out: list[object] = []
                    for column in range(sheet.ncols):
                        cell = sheet.cell(row, column)
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            try:
                                out.append(
                                    xlrd.xldate_as_datetime(cell.value, book.datemode)
                                )
                            except (ValueError, OverflowError):
                                out.append(cell.value)
                        else:
                            out.append(cell.value)
                    return out

                rows = [(row + 1, values(row)) for row in range(last)]
                for number, cells in _scan(rows, numbers).items():
                    result[(name, number)] = cells
        finally:
            book.release_resources()
        return result

    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        by_sheet: dict[str, set[int]] = defaultdict(set, wanted)
        for name, numbers in by_sheet.items():
            if name not in book.sheetnames:
                continue
            sheet = book[name]
            rows = (
                (number, list(values))
                for number, values in enumerate(
                    sheet.iter_rows(max_row=max(numbers), values_only=True), start=1
                )
            )
            for number, cells in _scan(list(rows), numbers, default_header).items():
                result[(name, number)] = cells
    finally:
        book.close()
    return result
