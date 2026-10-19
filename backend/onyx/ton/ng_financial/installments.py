"""The Controladoria's table of installment plans actually paid.

The NG books installment plans (parcelamentos) in a way that distorts the
DRE by millions, so the Controladoria keeps its own table of what was paid,
month by month. Its rows become records shaped like NG launches; at
normalization they replace the NG launches of the same natureza and month.

The table is read by its header, not by fixed columns: any sheet whose first
rows name a unit, an amount and a month (or a payment date) is read.
"""

import hashlib
import json
import re
import time
import unicodedata
from collections import Counter
from datetime import date, datetime
from io import BytesIO
from uuid import UUID
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from onyx.error_handling.error_codes import OnyxErrorCode
from onyx.error_handling.exceptions import OnyxError
from onyx.ton.ng_financial.models import (
    DiagnosticCode,
    DiagnosticLevel,
    ParsedImportResult,
    ParsedSourceRecordData,
    ParseDiagnostic,
    RowKind,
    SheetSummary,
)
from onyx.ton.ng_financial.parser import parse_date, parse_decimal
from onyx.ton.sources.models import SourceLocator

PROFILE_KEY = "installment_table"
PROFILE_VERSION = 1
SOURCE_KEY = "installments"
# Natureza used when the table has no account column.
DEFAULT_ACCOUNT = "PARCELAMENTOS"
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_SHEETS = 32
MAX_ROWS_PER_SHEET = 10_000
MAX_COLUMNS = 40
HEADER_SCAN_ROWS = 20
MAX_SECONDS = 30
# The header labels each field accepts, normalized (no accents, lower case).
FIELDS: dict[str, tuple[str, ...]] = {
    "unit": ("unidade", "unidade administrativa", "filial", "centro de custo"),
    "month": (
        "mes",
        "mes/ano",
        "mes de pagamento",
        "mes do pagamento",
        "competencia",
        "competencia de pagamento",
        "periodo",
        "referencia",
    ),
    "payment_date": ("data de pagamento", "data do pagamento", "pagamento", "data"),
    "amount": (
        "valor pago",
        "valor",
        "valor da parcela",
        "valor parcela",
        "valor pago da parcela",
    ),
    "account": (
        "conta ng",
        "conta",
        "codigo",
        "codigo estruturado",
        "natureza",
    ),
    "description": (
        "credor",
        "descricao",
        "historico",
        "orgao",
        "fornecedor",
        "parcelamento",
        "observacao",
        "observacoes",
    ),
    "document": ("documento", "numero", "n do parcelamento", "processo"),
    "installment": ("parcela", "n da parcela", "parcela n"),
}
COLUMN_MAP: dict[str, str] = {field: f"header:{field}" for field in FIELDS}
MONTH_NAMES = {
    "jan": 1,
    "fev": 2,
    "mar": 3,
    "abr": 4,
    "mai": 5,
    "jun": 6,
    "jul": 7,
    "ago": 8,
    "set": 9,
    "out": 10,
    "nov": 11,
    "dez": 12,
}


def _normal(value: object) -> str:
    if value is None:
        return ""
    # "Nº" loses its º before NFKD, which would turn it into an o.
    text = re.sub(r"[º°ª.:]", "", str(value).replace("\n", " "))
    text = unicodedata.normalize("NFKD", text)
    text = text.encode("ascii", "ignore").decode().casefold()
    return " ".join(text.split())


def _header(values: list[object]) -> dict[str, int] | None:
    found: dict[str, int] = {}
    for index, value in enumerate(values):
        label = _normal(value)
        if not label:
            continue
        for field, labels in FIELDS.items():
            if field not in found and label in labels:
                found[field] = index
                break
    has_period = "month" in found or "payment_date" in found
    if "unit" in found and "amount" in found and has_period:
        return found
    return None


def parse_month(value: object) -> date | None:
    """The first day of the month a cell names: a date, 01/2026, jan/26, 2026-01."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    text = _normal(value)
    match = re.fullmatch(r"(\d{1,2})[/-](\d{4})", text)
    if match:
        month, year = int(match.group(1)), int(match.group(2))
    elif match := re.fullmatch(r"(\d{4})[/-](\d{1,2})", text):
        year, month = int(match.group(1)), int(match.group(2))
    elif match := re.fullmatch(r"([a-z]+)[/ -]+(?:de )?(\d{2}|\d{4})", text):
        month = MONTH_NAMES.get(match.group(1)[:3], 0)
        year = int(match.group(2))
        year += 2000 if year < 100 else 0
    else:
        parsed = parse_date(value)
        return date(parsed.year, parsed.month, 1) if parsed else None
    if not 1 <= month <= 12 or not 2000 <= year <= 2100:
        raise ValueError("invalid month")
    return date(year, month, 1)


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().hour == 0 else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _digest(parts: object) -> str:
    payload = json.dumps(parts, sort_keys=True, ensure_ascii=False, default=str)
    return hashlib.sha256(payload.encode()).hexdigest()


class InstallmentTableParser:
    key = PROFILE_KEY

    def parse(  # noqa: C901 - one pass keeps the header in force for each row
        self, content: bytes, snapshot_id: UUID
    ) -> ParsedImportResult:
        if not content or len(content) > MAX_WORKBOOK_BYTES:
            raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Installment workbook size")
        start = time.monotonic()
        try:
            book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except (BadZipFile, KeyError, OSError, ValueError):
            raise OnyxError(
                OnyxErrorCode.INVALID_INPUT, "Invalid installment workbook"
            ) from None
        result = ParsedImportResult(
            profile_key=PROFILE_KEY,
            profile_version=PROFILE_VERSION,
            snapshot_id=snapshot_id,
            records=[],
            diagnostics=[],
            sheets=[],
        )

        def diagnostic(
            level: DiagnosticLevel,
            code: DiagnosticCode,
            sheet: str,
            row: int | None = None,
            column: int | None = None,
        ) -> None:
            result.diagnostics.append(
                ParseDiagnostic(
                    level=level,
                    code=code,
                    sheet_name=sheet,
                    row_number=row,
                    column=get_column_letter(column + 1)
                    if column is not None
                    else None,
                )
            )

        try:
            if len(book.worksheets) > MAX_SHEETS:
                raise OnyxError(OnyxErrorCode.INVALID_INPUT, "Installment sheet limit")
            for sheet in book.worksheets:
                kinds: Counter[RowKind] = Counter()
                duplicates: Counter[str] = Counter()
                current: dict[str, int] | None = None
                physical = 0
                for number, row in enumerate(
                    sheet.iter_rows(max_col=MAX_COLUMNS, values_only=True), start=1
                ):
                    if time.monotonic() - start > MAX_SECONDS:
                        raise OnyxError(
                            OnyxErrorCode.INVALID_INPUT, "Installment parser time limit"
                        )
                    physical = number
                    if number > MAX_ROWS_PER_SHEET:
                        raise OnyxError(
                            OnyxErrorCode.INVALID_INPUT, "Installment grid limit"
                        )
                    values = list(row)
                    mapping = _header(values)
                    if mapping is not None:
                        kinds[
                            RowKind.HEADER if current is None else RowKind.REPEATED_HEADER
                        ] += 1
                        current = mapping
                        continue
                    if current is None:
                        if number > HEADER_SCAN_ROWS:
                            break
                        continue
                    filled = [value for value in values if value not in (None, "")]
                    if not filled:
                        kinds[RowKind.BLANK] += 1
                        continue

                    def cell(field: str) -> object:
                        assert current is not None
                        index = current.get(field)
                        return values[index] if index is not None and index < len(values) else None

                    labels = " ".join(_normal(value) for value in filled[:3])
                    if "total" in labels.split() or labels.startswith(("total", "subtotal")):
                        kinds[RowKind.TOTAL] += 1
                        continue
                    try:
                        raw_amount = cell("amount")
                        amount = parse_decimal(
                            raw_amount if not isinstance(raw_amount, str) else raw_amount.strip()
                        )
                    except ValueError:
                        diagnostic(
                            DiagnosticLevel.ERROR,
                            DiagnosticCode.INVALID_AMOUNT,
                            sheet.title,
                            number,
                            current["amount"],
                        )
                        kinds[RowKind.UNKNOWN] += 1
                        continue
                    if amount is None:
                        diagnostic(
                            DiagnosticLevel.ERROR,
                            DiagnosticCode.MISSING_AMOUNTS,
                            sheet.title,
                            number,
                            current["amount"],
                        )
                        kinds[RowKind.UNKNOWN] += 1
                        continue
                    try:
                        paid = parse_date(cell("payment_date")) if "payment_date" in current else None
                    except ValueError:
                        diagnostic(
                            DiagnosticLevel.ERROR,
                            DiagnosticCode.INVALID_DATE,
                            sheet.title,
                            number,
                            current["payment_date"],
                        )
                        kinds[RowKind.UNKNOWN] += 1
                        continue
                    try:
                        month = parse_month(cell("month")) if "month" in current else None
                    except ValueError:
                        diagnostic(
                            DiagnosticLevel.ERROR,
                            DiagnosticCode.INVALID_DATE,
                            sheet.title,
                            number,
                            current["month"],
                        )
                        kinds[RowKind.UNKNOWN] += 1
                        continue
                    if month is None and paid is not None:
                        month = date(paid.year, paid.month, 1)
                    if month is None:
                        diagnostic(
                            DiagnosticLevel.ERROR,
                            DiagnosticCode.MISSING_DATE,
                            sheet.title,
                            number,
                            current.get("month", current.get("payment_date")),
                        )
                        kinds[RowKind.UNKNOWN] += 1
                        continue
                    unit = _text(cell("unit"))
                    if not unit:
                        diagnostic(
                            DiagnosticLevel.WARNING,
                            DiagnosticCode.UNIT_BLANK,
                            sheet.title,
                            number,
                            current["unit"],
                        )
                    account = _text(cell("account")) or DEFAULT_ACCOUNT
                    code, _, label = account.partition(" - ")
                    description = " · ".join(
                        part
                        for part in (
                            _text(cell("description")),
                            f"parcela {_text(cell('installment'))}"
                            if _text(cell("installment"))
                            else None,
                        )
                        if part
                    )
                    # What was paid leaves the cash: always an outflow.
                    signed = -abs(amount)
                    emission = paid if paid and (paid.year, paid.month) == (month.year, month.month) else month
                    fingerprint = _digest(
                        [
                            PROFILE_KEY,
                            PROFILE_VERSION,
                            code.strip(),
                            unit,
                            month.isoformat(),
                            emission.isoformat(),
                            str(signed),
                            description,
                            _text(cell("document")),
                        ]
                    )
                    duplicates[fingerprint] += 1
                    kinds[RowKind.DETAIL] += 1
                    result.records.append(
                        ParsedSourceRecordData(
                            locator=SourceLocator(
                                snapshot_id=snapshot_id,
                                sheet_name=sheet.title,
                                row_number=number,
                            ),
                            sheet_month=month.month,
                            account_code=code.strip(),
                            account_label=(label or code).strip(),
                            emission_date=emission,
                            administrative_unit=unit,
                            document_number=_text(cell("document")),
                            history=description or "Parcelamento pago (planilha)",
                            movement_amount=signed,
                            movement_retention_amount=None,
                            movement_net_amount=None,
                            installment_retention_amount=None,
                            installment_net_amount=signed,
                            interest_amount=None,
                            penalty_amount=None,
                            discount_amount=None,
                            expense_amount=None,
                            loss_amount=None,
                            other_deduction_amount=None,
                            final_amount=signed,
                            source_values={
                                get_column_letter(index + 1): _text(value)
                                for index, value in enumerate(values)
                                if value not in (None, "")
                            },
                            fingerprint=fingerprint,
                            duplicate_ordinal=duplicates[fingerprint],
                        )
                    )
                accepted = current is not None
                if not accepted:
                    diagnostic(
                        DiagnosticLevel.INFO,
                        DiagnosticCode.SHEET_OUT_OF_SCOPE,
                        sheet.title,
                    )
                result.sheets.append(
                    SheetSummary(
                        name=sheet.title,
                        month=None,
                        accepted=accepted,
                        physical_rows=physical,
                        classifications=dict(kinds),
                        row_spans={},
                    )
                )
        finally:
            book.close()
        if not result.records:
            raise OnyxError(
                OnyxErrorCode.INVALID_INPUT, "Unrecognized installment workbook"
            )
        return result

