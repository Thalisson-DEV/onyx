"""Synthetic DRE workbook: formulas over Base, cached values and recalculation."""

import datetime
import io
import re
from decimal import Decimal

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from onyx.ton.dre.models import DreLineDefinition, DreLineType, DreOperation
from onyx.ton.dre.xlsx_export import (
    BaseEntry,
    DreWorkbookInput,
    Premise,
    ScopeBlock,
    build_workbook,
    expected_values,
)

CELL = re.compile(r"\$?([A-Z]+)\$?(\d+)")


def _at(sheet: Worksheet, reference: str) -> object:
    cell = sheet[reference]
    assert isinstance(cell, Cell)
    return cell.value


def _lines() -> list[DreLineDefinition]:
    return [
        DreLineDefinition(
            code="g1",
            label="Receita",
            position=0,
            line_type=DreLineType.SUBTOTAL,
            operation=DreOperation.SUM_CHILDREN,
        ),
        DreLineDefinition(
            code="n1.01",
            label="Serviço",
            position=1,
            parent_code="g1",
            line_type=DreLineType.SOURCE_SUM,
        ),
        DreLineDefinition(
            code="n1.02",
            label="Outras receitas",
            position=2,
            parent_code="g1",
            line_type=DreLineType.SOURCE_SUM,
        ),
        DreLineDefinition(
            code="n2.01",
            label="Custo",
            position=3,
            line_type=DreLineType.SOURCE_SUM,
        ),
        DreLineDefinition(
            code="r1",
            label="Resultado",
            position=4,
            line_type=DreLineType.RESULT,
            operation=DreOperation.SUBTRACT,
            operands=["g1", "n2.01"],
        ),
        DreLineDefinition(
            code="r2",
            label="Resultado + receita",
            position=5,
            line_type=DreLineType.RESULT,
            operation=DreOperation.SUM_LINES,
            operands=["r1", "g1"],
        ),
        DreLineDefinition(
            code="m1",
            label="Margem %",
            position=6,
            line_type=DreLineType.PERCENTAGE,
            operation=DreOperation.RATIO,
            operands=["r1", "g1"],
        ),
    ]


def _entry(
    month: int,
    unit: str | None,
    line: str,
    amount: str,
    row: int,
    history: str = "Lançamento sintético",
) -> BaseEntry:
    return BaseEntry(
        competence=datetime.date(2026, month, 1),
        unit_code=unit,
        unit_name=f"Unidade {unit}" if unit else None,
        line_code=line,
        amount=Decimal(amount),
        record_date=datetime.date(2026, month, 10),
        nature="Natureza sintética",
        ng_account_code="9.9.9",
        ng_account_label="Conta sintética",
        document=f"SYN-{row}",
        history=history,
        amount_basis="MOVEMENT",
        review_status="ACCEPTED",
        source_file="synthetic.xlsx",
        sheet_name="Jan",
        row_number=row,
    )


def _input() -> DreWorkbookInput:
    entries = [
        _entry(1, "000001", "n1.01", "1000.10", 1),
        _entry(1, "000001", "n2.01", "400.05", 2),
        _entry(1, "000002", "n1.01", "250.25", 3),
        _entry(1, None, "n1.02", "10.00", 4),
        _entry(2, "000001", "n1.01", "800.00", 5, history="=HYPERLINK(1)"),
        _entry(2, "000001", "n1.02", "33.33", 6),
        _entry(2, "000001", "n2.01", "900.40", 7),
        _entry(3, "000001", "n1.01", "120.00", 8),
        _entry(3, "000002", "n2.01", "75.75", 9),
    ]
    return DreWorkbookInput(
        structure_label="Synthetic gerencial",
        structure_version=2,
        lines=_lines(),
        year=2026,
        last_month=3,
        scopes=[
            ScopeBlock(None, "Consolidado", frozenset({1, 2, 3})),
            ScopeBlock("000001", "Unidade 000001", frozenset({1, 2, 3})),
            ScopeBlock("000002", "Unidade 000002", frozenset({1})),
        ],
        entries=entries,
        header=[("Período", "jan a mar/2026")],
        premises=[Premise("Orçado", "Zerado", "Só Realizado")],
    )


def _number(value: object) -> Decimal:
    if value in (None, ""):
        return Decimal(0)
    assert isinstance(value, (int, float))
    return Decimal(str(value))


class _Recalculator:
    """Evaluates the formula shapes the export writes, from the Base rows."""

    def __init__(self, workbook: openpyxl.Workbook) -> None:
        self.dre: Worksheet = workbook["DRE"]
        base: Worksheet = workbook["Base"]
        rows = list(base.iter_rows(values_only=True))
        header = list(rows[0])
        self.base = [dict(zip(header, row, strict=True)) for row in rows[1:]]
        self.cache: dict[str, Decimal | None] = {}

    def value(self, reference: str) -> Decimal | None:
        reference = reference.replace("$", "")
        if reference not in self.cache:
            raw = _at(self.dre, reference)
            if isinstance(raw, str) and raw.startswith("="):
                self.cache[reference] = self.formula(raw[1:])
            else:
                self.cache[reference] = _number(raw)
        return self.cache[reference]

    def raw(self, reference: str) -> object:
        return _at(self.dre, reference.replace("$", ""))

    def formula(self, text: str) -> Decimal | None:
        if text.startswith("SUMIFS("):
            parts = text[len("SUMIFS(") : -1].split(",")
            assert parts[0] == "Base[Valor]"
            criteria = [
                (parts[index][len("Base[") : -1], self.raw(parts[index + 1]))
                for index in range(1, len(parts), 2)
            ]
            return sum(
                (
                    _number(row["Valor"])
                    for row in self.base
                    if all(row[column] == wanted for column, wanted in criteria)
                ),
                Decimal(0),
            )
        if text.startswith("SUM("):
            first, last = text[4:-1].split(":")
            start = CELL.fullmatch(first)
            end = CELL.fullmatch(last)
            assert start and end and start.group(2) == end.group(2)
            return sum(
                (
                    self.value(f"{chr(column)}{start.group(2)}") or Decimal(0)
                    for column in range(ord(start.group(1)), ord(end.group(1)) + 1)
                ),
                Decimal(0),
            )
        ratio = re.fullmatch(r'IF\((\w+)=0,"",(\w+)/(\w+)\*100\)', text)
        if ratio:
            denominator = self.value(ratio.group(1))
            numerator = self.value(ratio.group(2))
            if not denominator or numerator is None:
                return None
            return numerator / denominator * 100
        if "-" in text:
            left, right = text.split("-")
            a, b = self.value(left), self.value(right)
            assert a is not None and b is not None
            return a - b
        values = [self.value(part) for part in text.split("+")]
        assert all(item is not None for item in values)
        return sum((item for item in values if item is not None), Decimal(0))


def _scope_rows(sheet: Worksheet) -> dict[str | None, dict[str, int]]:
    """Row of each line code per block, keyed by the block's unit code."""
    blocks: dict[str | None, dict[str, int]] = {}
    current: str | None = None
    for row in range(1, sheet.max_row + 1):
        first = sheet.cell(row, 1).value
        if sheet.cell(row + 2, 1).value == "Código" and first:
            current = None if first == "Consolidado" else str(first)
            blocks[current] = {}
        elif current in blocks and first and first != "Código":
            blocks[current][str(first)] = row
    return blocks


def test_workbook_formulas_cache_and_recalculation_match() -> None:
    data = _input()
    content = build_workbook(data)
    formulas = openpyxl.load_workbook(io.BytesIO(content))
    cached = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    assert formulas.sheetnames == ["DRE", "Base", "Premissas"]

    expected = expected_values(data)
    # Consolidated includes the entry without unit; units do not.
    assert expected[None][(1, "g1")] == Decimal("1260.35")
    assert expected["000001"][(None, "r1")] == Decimal("652.98")
    assert expected["000002"][(3, "m1")] is None

    recalculator = _Recalculator(formulas)
    sheet = formulas["DRE"]
    blocks = _scope_rows(sheet)
    assert set(blocks) == {None, "000001", "000002"}
    for unit, rows in blocks.items():
        assert set(rows) == {line.code for line in data.lines}
        for code, row in rows.items():
            for column, month in zip("CDEF", (1, 2, 3, None), strict=True):
                cell = f"{column}{row}"
                formula = _at(sheet, cell)
                assert isinstance(formula, str) and formula.startswith("=")
                if code.startswith("n") and month is not None:
                    assert formula.startswith("=SUMIFS(Base[Valor],")
                    assert ("Base[Unidade]" in formula) == (unit is not None)
                want = expected[unit][(month, code)]
                cached_value = _at(cached["DRE"], cell)
                recalculated = recalculator.value(cell)
                if want is None:
                    assert cached_value in (None, "")
                    assert recalculated is None
                    continue
                assert recalculated is not None
                assert recalculated.quantize(Decimal("0.01")) == want.quantize(
                    Decimal("0.01")
                )
                assert _number(cached_value).quantize(Decimal("0.01")) == (
                    want.quantize(Decimal("0.01"))
                )

    base = formulas["Base"]
    assert base.tables["Base"].ref == "A1:P10"
    histories = [row[11] for row in base.iter_rows(min_row=2, values_only=True)]
    # NG text that looks like a formula stays text.
    assert "=HYPERLINK(1)" in histories
    assert all(
        not isinstance(cell.value, str)
        or not cell.value.startswith("=")
        or cell.column == 12
        for row in base.iter_rows(min_row=2)
        for cell in row
    )
    premises = [row for row in formulas["Premissas"].iter_rows(values_only=True)]
    assert ("Orçado", "Zerado", "Só Realizado") in premises
    statuses = [
        sheet.cell(row, column).value
        for row in range(1, sheet.max_row + 1)
        for column in (3, 4, 5)
        if sheet.cell(row, 2).value == "Situação da DRE no TON"
    ]
    assert statuses.count("Não pronta") == 2


def test_empty_base_still_builds() -> None:
    data = _input()
    empty = DreWorkbookInput(
        structure_label=data.structure_label,
        structure_version=data.structure_version,
        lines=data.lines,
        year=2026,
        last_month=1,
        scopes=[ScopeBlock(None, "Consolidado", frozenset({1}))],
        entries=[],
        header=[],
    )
    workbook = openpyxl.load_workbook(io.BytesIO(build_workbook(empty)), data_only=True)
    assert workbook["Base"].tables["Base"].ref == "A1:P2"
