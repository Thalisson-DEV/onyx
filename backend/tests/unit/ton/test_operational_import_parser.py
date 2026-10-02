"""Synthetic source workbooks for DATA-004A/B."""

from datetime import date, datetime
from io import BytesIO
from typing import cast
from uuid import UUID
from zipfile import ZipFile

import pytest
import xlwt
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from onyx.error_handling.exceptions import OnyxError
from onyx.ton.operational_import import parser as parser_module
from onyx.ton.operational_import.models import SheetClass
from onyx.ton.operational_import.parser import (
    BILLING_KEY,
    BUDGET_ANNUAL_KEY,
    BUDGET_TERM_KEY,
    BillingParser,
    BudgetParser,
    identify_profile,
)
from onyx.ton.sources.models import SourceFormat

SNAPSHOT = UUID(int=2)
HEADER = (
    "Tomador",
    "Nº da Nota Fiscal",
    "Data de Emissão",
    "Valor Serviço",
    "ISS Retido",
    "Total Retido",
    "Líquido Após o desc.",
    "Comp",
)


def billing_book(
    rows: list[tuple[object, ...]],
    *,
    header: tuple[str, ...] = HEADER,
) -> bytes:
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Export")
    sheet.write(0, 0, "Metadata")
    sheet.write(1, 0, "Report")
    for column, value in enumerate(header):
        sheet.write(2, column, value)
    style = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for row_number, values in enumerate(rows, 3):
        for column, value in enumerate(values):
            if value is None:
                continue
            sheet.write(
                row_number,
                column,
                value,
                style if isinstance(value, datetime) else xlwt.Style.default_style,
            )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def invoice(
    number: object = 101,
    gross: object = 100.25,
    retained: object = 0,
    competence: object = "01/2026",
    emitted: datetime = datetime(2026, 1, 10),
) -> tuple[object, ...]:
    return (
        "Synthetic payer",
        number,
        emitted,
        gross,
        retained,
        retained,
        0,
        competence,
    )


def budget_book(
    key: str = BUDGET_ANNUAL_KEY,
    *,
    formula_cache: bool = True,
    malformed: bool = False,
    merged: bool = False,
) -> bytes:
    workbook = Workbook()
    sheet = cast(Worksheet, workbook.active)
    sheet.title = "DOTAÇÃO"
    sheet["A1"] = "Synthetic contract"
    sheet["A2"] = (
        "Valor do Contrato Anual"
        if key == BUDGET_ANNUAL_KEY
        else "Valor do Contrato (30 meses)"
    )
    sheet["A3"] = "Valor do Contrato Mensal"
    sheet.cell(2, 3).value = 120 if key == BUDGET_ANNUAL_KEY else "=C3*30"
    sheet.cell(3, 3).value = "=C2/12" if key == BUDGET_ANNUAL_KEY else 10
    sheet["A8"] = "ITEM"
    sheet["B8"] = "DESCRIÇÃO"
    sheet["C8"] = "VALOR TOTAL"
    sheet["D8"] = "%"
    sheet["A10"] = "1"
    sheet["B10"] = "COMPOSIÇÃO DE CUSTOS SINTÉTICO"
    sheet["A11"] = "1.1"
    sheet["B11"] = "Section"
    sheet["A12"] = "1.1.1"
    sheet["B12"] = "Synthetic detail"
    sheet["C12"] = "=5+5"
    sheet["A13"] = "1.1.2"
    sheet["B13"] = "Second detail"
    sheet.cell(13, 3).value = "bad" if malformed else 0
    sheet["A35"] = "2"
    sheet["B35"] = "COMPOSIÇÃO DA PROPOSTA FINANCEIRA"
    sheet["C35"] = "=C12+C13"
    if merged:
        sheet.merge_cells("B12:B13")
    summary = workbook.create_sheet("SÍNTESE DOS CUSTOS")
    summary["A1"] = "Summary"
    output = BytesIO()
    workbook.save(output)
    if not formula_cache:
        return output.getvalue()
    patched = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as source, ZipFile(patched, "w") as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                payload = payload.replace(
                    b'<c r="C12"><f>5+5</f><v></v></c>',
                    b'<c r="C12"><f>5+5</f><v>10</v></c>',
                )
            target.writestr(item, payload)
    return patched.getvalue()


def test_billing_structural_header_types_and_lineage() -> None:
    content = billing_book([invoice(), HEADER, invoice(102, "R$ 1.234,50", None)])
    result = BillingParser().parse(content, SNAPSHOT)
    assert result.profile_key == BILLING_KEY
    assert len(result.records) == 2
    assert result.records[0].locator.sheet_name == "Export"
    assert result.records[0].locator.row_number == 4
    assert result.records[1].locator.row_number == 6
    assert result.records[0].competence == date(2026, 1, 1)
    assert str(result.records[1].amount) == "1234.50"
    assert result.records[0].numeric_values["iss_retained"] == 0
    assert result.records[1].numeric_values["iss_retained"] is None
    assert identify_profile(content, SourceFormat.XLS) == BILLING_KEY


def test_billing_rejections_duplicates_and_determinism() -> None:
    content = billing_book(
        [
            invoice(),
            invoice(),
            invoice(gross="invalid"),
            invoice(number="-"),
            invoice(competence="invalid"),
        ]
    )
    first = BillingParser().parse(content, SNAPSHOT)
    second = BillingParser().parse(content, SNAPSHOT)
    assert len(first.records) == 2
    assert [record.duplicate_ordinal for record in first.records] == [1, 2]
    assert [record.fingerprint for record in first.records] == [
        record.fingerprint for record in second.records
    ]
    assert {item.code for item in first.diagnostics} == {
        "INVALID_AMOUNT",
        "INVALID_INVOICE",
        "INVALID_COMPETENCE",
    }


def test_billing_wrong_structure_and_format() -> None:
    with pytest.raises(OnyxError):
        BillingParser().parse(billing_book([], header=("Name", "Value")), SNAPSHOT)
    with pytest.raises(OnyxError):
        BillingParser().parse(b"not an XLS", SNAPSHOT)
    with pytest.raises(OnyxError):
        identify_profile(billing_book([invoice()]), SourceFormat.XLSX)


@pytest.mark.parametrize("key", [BUDGET_ANNUAL_KEY, BUDGET_TERM_KEY])
def test_budget_families_detail_formula_and_summary(key: str) -> None:
    content = budget_book(key)
    result = BudgetParser(key).parse(content, SNAPSHOT)
    assert identify_profile(content, SourceFormat.XLSX) == key
    assert len(result.records) == 2
    assert [item.locator.row_number for item in result.records] == [12, 13]
    assert result.records[0].formula_cached
    assert result.records[0].amount == 10
    assert result.records[1].amount == 0
    assert result.records[0].period_basis == "MONTHLY_CONTRACT"
    assert result.sheets[0].aggregates_excluded == 2
    assert result.sheets[0].rows_inspected == result.sheets[0].physical_rows
    assert result.sheets[1].classification == SheetClass.SUMMARY
    assert result.sheets[1].rows_inspected == 0
    assert all(record.locator.sheet_name == "DOTAÇÃO" for record in result.records)


def test_budget_partial_formula_cache_and_merged_handling() -> None:
    uncached = BudgetParser(BUDGET_ANNUAL_KEY).parse(
        budget_book(formula_cache=False), SNAPSHOT
    )
    assert len(uncached.records) == 1
    assert uncached.error_count == 1
    assert uncached.diagnostics[0].code == "FORMULA_CACHE_MISSING"
    malformed = BudgetParser(BUDGET_ANNUAL_KEY).parse(
        budget_book(malformed=True), SNAPSHOT
    )
    assert malformed.error_count == 1
    assert malformed.diagnostics[0].code == "INVALID_AMOUNT"
    merged = BudgetParser(BUDGET_ANNUAL_KEY).parse(budget_book(merged=True), SNAPSHOT)
    assert len(merged.records) == 1
    assert merged.diagnostics[0].code == "MERGED_DETAIL_UNSUPPORTED"


def test_budget_wrong_family_and_unsupported_bytes() -> None:
    with pytest.raises(OnyxError):
        BudgetParser(BUDGET_TERM_KEY).parse(budget_book(), SNAPSHOT)
    with pytest.raises(OnyxError):
        BudgetParser(BUDGET_ANNUAL_KEY).parse(b"not XLSX", SNAPSHOT)


def test_bounded_workbooks(monkeypatch: pytest.MonkeyPatch) -> None:
    billing = billing_book([invoice()])
    budget = budget_book()
    monkeypatch.setattr(parser_module, "MAX_BYTES", 100)
    with pytest.raises(OnyxError):
        BillingParser().parse(billing, SNAPSHOT)
    with pytest.raises(OnyxError):
        BudgetParser(BUDGET_ANNUAL_KEY).parse(budget, SNAPSHOT)
    monkeypatch.setattr(parser_module, "MAX_BYTES", 10 * 1024 * 1024)
    monkeypatch.setattr(parser_module, "MAX_ROWS_PER_SHEET", 3)
    with pytest.raises(OnyxError):
        BillingParser().parse(billing, SNAPSHOT)
    with pytest.raises(OnyxError):
        BudgetParser(BUDGET_ANNUAL_KEY).parse(budget, SNAPSHOT)
